"""
Excel della revisione con indicazione precompilata: per ogni ponte con posizione proposta l'indicazione
(SI / probabile SI / da guardare), il motivo, cosa guardare sull'ortofoto e la colonna Conferma da compilare.

Uso: python indicazioni_revisione_excel.py <indicazioni_revisione.csv> <revisione.csv>
        <sezioni_su_tracciato_kml.csv> <Database.xlsx> <out.xlsx>
"""
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ind, rev, trk, dbx, out = sys.argv[1:6]
I = pd.read_csv(ind)
R = pd.read_csv(rev).set_index('ID')
K = pd.read_csv(trk).set_index('ID')
db = pd.read_excel(dbx).rename(columns={'Unnamed: 0': 'ID'})
db = db[pd.to_numeric(db.ID, errors='coerce').notna()]
db['ID'] = db.ID.astype(int)
db = db.set_index('ID')


def guarda(r):
    rel = str(r.Relazione)
    if r.Indicazione.startswith('SI (confluenza'):
        return "Verificare che l'affluente entri nell'alveo sotto il ponte o subito accanto."
    if r.Indicazione == 'SI':
        return ''
    if rel.startswith('stesso corso'):
        aff = r['Area affluente maggiore (km²)'] if pd.notna(r['Area affluente maggiore (km²)']) else 0
        if aff <= 0.1 * r.A_db_km2 and r['Dist proposto-tracciato (m)'] <= 2 * r['Soglia sotto il ponte (m)']:
            return ('Probabile SI: stesso corso, nessun affluente importante in mezzo, punto appena fuori '
                    "dall'impronta del ponte. Basta controllare che il ponte sia su questo corso.")
        return (f"Lo spostamento scavalca un affluente di {aff:g} km²: guardare se il ponte e' prima o dopo la "
                "confluenza (se l'affluente entra a valle del ponte, NO).")
    if not r['Tracciato attraversa corso proposto'] and r['Tracciato attraversa corso attuale']:
        return "Il tracciato del ponte attraversa il corso attuale, non quello proposto: probabile NO, verificare."
    return "Guardare sull'ortofoto quale dei due corsi attraversa il ponte."


def nome(i):
    n = str(db.loc[i, 'Nome opera']).strip()
    return '' if n in ('-', 'nan') else n


T = pd.DataFrame({
    'ID': I.ID, 'Nome opera': [nome(i) for i in I.ID], 'Origine proposta': I['Origine proposta'],
    'Area DB (km²)': I.A_db_km2, 'Area attuale (km²)': [R.A_attuale_km2.get(i) for i in I.ID],
    'Area proposta (km²)': [K.A_sezione_tracciato_km2[i] if i in K.index else R.A_proposta_km2.get(i) for i in I.ID],
    'Indicazione': I.Indicazione, 'Motivo': I.Motivo, 'Cosa guardare': [guarda(r) for _, r in I.iterrows()],
    'Relazione tra i punti': I.Relazione, 'Distanza lungo asta (m)': I['Distanza lungo asta (m)'],
    'Affluenti scavalcati': I['Affluenti scavalcati'], 'Area affluente maggiore (km²)': I['Area affluente maggiore (km²)'],
    'Dist proposto-tracciato (m)': I['Dist proposto-tracciato (m)'], 'Soglia sotto il ponte (m)': I['Soglia sotto il ponte (m)'],
    'Tracciato attraversa corso proposto': I['Tracciato attraversa corso proposto'].map({True: 'si', False: 'no'}),
    'Tracciato attraversa corso attuale': I['Tracciato attraversa corso attuale'].map({True: 'si', False: 'no'}),
    'Conferma (SI/NO)': '', 'Note': ''})
T.loc[T['Cosa guardare'].str.startswith('Probabile SI'), 'Indicazione'] = 'probabile SI'
ordine = {'SI': 0, 'SI (confluenza al ponte)': 1, 'probabile SI': 2, 'da guardare': 3}
T = T.sort_values(['Indicazione', 'ID'], key=lambda s: s.map(ordine) if s.name == 'Indicazione' else s)
with pd.ExcelWriter(out, engine='openpyxl') as w:
    T.to_excel(w, sheet_name='Indicazioni', index=False)
    pd.DataFrame({'Criteri': [
        "1) Stesso corso d'acqua? Il punto proposto deve stare sull'asta del DEM che rappresenta il corso attraversato dal ponte.",
        "2) Sotto il ponte? Distanza del punto proposto dal tracciato KML <= max(lunghezza ponte, luce, alveo)/2 + 20 m.",
        "3) Confluenze? Se lo spostamento scavalca un affluente, l'affluente deve entrare nell'alveo sotto il ponte o subito accanto.",
        "SI = criteri soddisfatti; SI (confluenza al ponte) = verificare solo la confluenza; probabile SI = stesso corso appena fuori "
        "dall'impronta del ponte; da guardare = serve l'ortofoto.",
        "Le aree sono quelle della prova (stessa catena del codice); la conferma e' sempre dell'utente."]}).to_excel(
        w, sheet_name='Criteri', index=False)
wb = load_workbook(out)
ws = wb['Indicazioni']
n = ws.max_row
dv = DataValidation(type='list', formula1='"SI,NO"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f'R2:R{n}')
for r in range(2, n + 1):
    ws[f'R{r}'].fill = PatternFill('solid', fgColor='FFF2CC')
    for c in ws[r]:
        c.alignment = Alignment(wrap_text=True, vertical='top')
    for c in 'DEF':
        ws[f'{c}{r}'].number_format = '0.000'
colori = {'SI': 'C6EFCE', 'SI (confluenza al ponte)': 'E2EFDA', 'probabile SI': 'DDEBF7', 'da guardare': 'FCE4D6'}
for v, col in colori.items():
    ws.conditional_formatting.add(f'G2:G{n}', FormulaRule(formula=[f'G2="{v}"'], fill=PatternFill('solid', fgColor=col)))
for c, wdt in zip('ABCDEFGHIJKLMNOPQRS', [6, 22, 18, 10, 10, 10, 16, 45, 50, 30, 10, 10, 11, 11, 11, 11, 11, 11, 25]):
    ws.column_dimensions[c].width = wdt
ws.freeze_panes = 'C2'
for w_ in wb.worksheets:
    for c in w_[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDDDDD')
        c.alignment = Alignment(wrap_text=True, vertical='center')
wb['Criteri'].column_dimensions['A'].width = 140
wb.save(out)
print(T[['ID', 'Indicazione', 'Cosa guardare']].to_string(max_colwidth=110))
