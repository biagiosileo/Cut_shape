"""
Excel di confronto tra area calcolata dal codice e SUPERFICIE BACINO del database (2 fogli, con formule):
  1 Statistiche      soglie di tolleranza (modificabili) e conteggi per classe con COUNTIF
  2 Confronto aree   ID, nome, area database, area calcolata, differenza (km²), differenza %,
                     differenza % assoluta, classe, criterio; quasi compatibili in celeste, non compatibili in giallo.
Classi: compatibile |diff %| <= 15%; quasi compatibile |diff %| <= 25%, oppure bacino piccolo con
|diff| <= 0,1 km² (soglia del reticolo) e |diff %| <= 50%; non compatibile altrimenti.

Uso: python excel_confronto_aree_finale.py <parametri_morfologici.csv> <Database.xlsx> <out.xlsx>
"""
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

pcsv, dbx, out = sys.argv[1:4]
U = pd.read_csv(pcsv)
U['ID'] = U.ID.astype(int)
assert U.ID.tolist() == list(range(1, 213)), 'ID non 1-212 in ordine'
db = pd.read_excel(dbx).rename(columns={'Unnamed: 0': 'ID'})
db = db[pd.to_numeric(db.ID, errors='coerce').notna()]
db['ID'] = db.ID.astype(int)
db = db.set_index('ID')
adb = pd.to_numeric(db['SUPERFICIE BACINO (km2)'], errors='coerce')

CELESTE = PatternFill('solid', fgColor='BDD7EE')
GIALLO = PatternFill('solid', fgColor='FFF2CC')
GRIGIO = PatternFill('solid', fgColor='D9D9D9')
BLU = Font(color='0000FF')
GRASS = Font(bold=True)
sottile = Side(style='thin', color='999999')
bordo = Border(top=sottile, bottom=sottile, left=sottile, right=sottile)

wb = Workbook()
st = wb.active
st.title = 'Statistiche'
ca = wb.create_sheet('Confronto aree')

# --- foglio 2: confronto ---
n = len(U)
ult = n + 1
intest = ['ID', 'Nome opera', 'Area database (km²)', 'Area calcolata (km²)', 'Differenza (km²)',
          'Differenza %', 'Differenza % assoluta', 'Classe', 'Criterio']
ca.append(intest)
for _, r in U.iterrows():
    i = r.ID
    nome = str(db.loc[i, 'Nome opera']).strip() if i in db.index else ''
    a_db = adb.get(i)
    ca.append([i, '' if nome in ('-', 'nan') else nome, None if pd.isna(a_db) else float(a_db), float(r.A_km2)])
for k in range(2, ult + 1):
    ca[f'E{k}'] = f'=IF(C{k}="","",D{k}-C{k})'
    ca[f'F{k}'] = f'=IF(C{k}="","",D{k}/C{k}-1)'
    ca[f'G{k}'] = f'=IF(C{k}="","",ABS(F{k}))'
    ca[f'H{k}'] = (f'=IF(C{k}="","senza area DB",IF(G{k}<=Statistiche!$C$4,"compatibile",'
                   f'IF(OR(G{k}<=Statistiche!$C$5,AND(ABS(E{k})<=Statistiche!$C$6,G{k}<=Statistiche!$C$7)),'
                   f'"quasi compatibile","non compatibile")))')
    ca[f'I{k}'] = (f'=IF(H{k}="quasi compatibile",IF(G{k}<=Statistiche!$C$5,"differenza % entro la tolleranza",'
                   f'"bacino piccolo: differenza assoluta entro la soglia"),"")')
    for c in 'CD':
        ca[f'{c}{k}'].number_format = '0.000'
        ca[f'{c}{k}'].font = BLU
    ca[f'E{k}'].number_format = '+0.000;-0.000;0.000'
    for c in 'FG':
        ca[f'{c}{k}'].number_format = '+0.0%;-0.0%;0.0%' if c == 'F' else '0.0%'
    for c in 'ABCDEFGHI':
        ca[f'{c}{k}'].border = bordo
ca.conditional_formatting.add(f'A2:I{ult}', FormulaRule(formula=['$H2="quasi compatibile"'], fill=CELESTE))
ca.conditional_formatting.add(f'A2:I{ult}', FormulaRule(formula=['$H2="non compatibile"'], fill=GIALLO))
ca.conditional_formatting.add(f'A2:I{ult}', FormulaRule(formula=['$H2="senza area DB"'], fill=GRIGIO))
for c, w in zip('ABCDEFGHI', [6, 44, 14, 14, 14, 13, 13, 18, 44]):
    ca.column_dimensions[c].width = w
ca.freeze_panes = 'C2'
ca.auto_filter.ref = f'A1:I{ult}'

# --- foglio 1: statistiche ---
st['A1'] = 'Confronto aree: area calcolata dal codice e SUPERFICIE BACINO del database'
st['A1'].font = Font(bold=True, size=12)
st['A3'], st['B3'] = 'Tolleranze', 'Classe'
st['A4'], st['B4'], st['C4'] = '|differenza %| ≤', 'compatibile', 0.15
st['A5'], st['B5'], st['C5'] = '|differenza %| ≤', 'quasi compatibile (oltre la soglia precedente)', 0.25
st['A6'], st['B6'], st['C6'] = '|differenza| (km²) ≤', 'quasi compatibile anche se: bacino piccolo (soglia del reticolo)', 0.1
st['A7'], st['B7'], st['C7'] = 'e |differenza %| ≤', 'limite per il criterio del bacino piccolo', 0.5
st['A8'], st['B8'] = 'altrimenti', 'non compatibile'
for c in ('C4', 'C5', 'C7'):
    st[c].number_format = '0%'
st['C6'].number_format = '0.00'
st['C4'].font = st['C5'].font = st['C6'].font = st['C7'].font = BLU
st['D4'] = 'valori modificabili: le classi si aggiornano'

st['A10'], st['B10'], st['C10'], st['D10'], st['E10'] = ('Classe', 'Colore', 'Numero ponti', '% sul totale',
                                                   '% sui ponti con area DB')
classi = [('compatibile', 'bianco', None), ('quasi compatibile', 'celeste', CELESTE),
          ('non compatibile', 'giallo', GIALLO), ('senza area DB', 'grigio', GRIGIO)]
for k, (cl, col, fill) in enumerate(classi, start=11):
    st[f'A{k}'], st[f'B{k}'] = cl, col
    st[f'C{k}'] = f"=COUNTIF('Confronto aree'!$H$2:$H${ult},A{k})"
    st[f'D{k}'] = f'=C{k}/$C$15'
    st[f'E{k}'] = f'=IF(A{k}="senza area DB","",C{k}/$C$16)'
    if fill:
        for c in 'ABCDE':
            st[f'{c}{k}'].fill = fill
st['A15'], st['C15'] = 'Totale ponti', '=SUM(C11:C14)'
st['D15'] = '=C15/C15'
st['A16'], st['C16'] = 'Ponti con area nel database', '=C15-C14'
st['A17'], st['C17'] = 'Fuori tolleranza (quasi + non compatibili)', '=C12+C13'
st['E17'] = '=C17/C16'
st['A18'], st['C18'] = 'di cui quasi compatibili per il criterio del bacino piccolo', (
    f"=COUNTIF('Confronto aree'!$I$2:$I${ult},\"bacino piccolo*\")")
st['A19'], st['C19'] = 'Differenza % mediana (ponti con area DB)', f"=MEDIAN('Confronto aree'!$F$2:$F${ult})"
st['A20'], st['C20'] = "Differenza % assoluta mediana", f"=MEDIAN('Confronto aree'!$G$2:$G${ult})"
st['A21'], st['C21'] = 'Differenza % assoluta media (solo compatibili)', (
    f"=AVERAGEIF('Confronto aree'!$H$2:$H${ult},\"compatibile\",'Confronto aree'!$G$2:$G${ult})")
for c in ('D11', 'D12', 'D13', 'D14', 'D15', 'E11', 'E12', 'E13', 'E17'):
    st[c].number_format = '0.0%'
for c in ('C19',):
    st[c].number_format = '+0.0%;-0.0%;0.0%'
for c in ('C20', 'C21'):
    st[c].number_format = '0.0%'
for r in (3, 10):
    for c in 'ABCDE':
        st[f'{c}{r}'].font = GRASS
for r in (15, 16, 17):
    st[f'A{r}'].font = GRASS
for r in range(10, 18):
    for c in 'ABCDE':
        st[f'{c}{r}'].border = bordo
st['A23'] = ("Differenza = area calcolata - area database; differenza % = area calcolata / area database - 1. "
             "Le classi del foglio 'Confronto aree' sono calcolate con le tolleranze in C4-C7.")
st['A23'].alignment = Alignment(wrap_text=True, vertical='top')
st.merge_cells('A23:E24')
for c, w in zip('ABCDE', [48, 58, 14, 13, 22]):
    st.column_dimensions[c].width = w
for ws in (ca,):
    for c in ws[1]:
        c.font = GRASS
        c.fill = PatternFill('solid', fgColor='DDDDDD')
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = bordo
wb.save(out)
print('ok', out)
