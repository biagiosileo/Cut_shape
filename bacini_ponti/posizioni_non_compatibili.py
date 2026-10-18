"""
Shapefile delle posizioni attuali e nuove dei ponti non compatibili (ID del database),
con stile QGIS (.qml) per categoria, e Excel con la giustificazione per i ponti in cui
l'errore e' attribuito al database.

Categorie della posizione nuova (da ricerca_sezione_compatibile.py):
  Spostamento <=50 m          sezione spostata sul punto compatibile (verde)
  Spostamento 50-200 m        sezione spostata, da verificare su cartografia (arancio)
  Area DB sotto soglia        sezione sul fosso con area compatibile, fuori reticolo (azzurro)
  Probabile errore database   sezione NON spostata: nessuna sezione compatibile entro 200 m (rosso)

Per i ponti con errore nel database si confrontano area DB e area calcolata con l'area
stimata dalla larghezza dell'alveo inciso (relazione ricavata dai ponti compatibili), per
capire se l'errore e' nel valore dell'area o nelle coordinate/posizione del ponte.

Uso: python posizioni_non_compatibili.py <Ricerca_sezione_compatibile.xlsx> <ricerca_sezione_compatibile.gpkg>
        <risultati_5m.gpkg> <Ponti_non_compatibili.xlsx> <Database.xlsx> <out_dir>
"""
import os
import sys

import numpy as np
import pandas as pd
import geopandas as gpd
from shapely.geometry import LineString

xr, gr, rgpkg, xnc, dbx, out = sys.argv[1:7]
os.makedirs(out, exist_ok=True)
R = pd.read_excel(xr).set_index('ID')
C = gpd.read_file(gr, layer='sezioni_compatibili')
C = C[C.tipo == 'compatibile più vicina'].set_index('ID')
S = gpd.read_file(rgpkg, layer='sezioni_chiusura').set_index('ID')
NC = pd.read_excel(xnc).set_index('ID')
db = pd.read_excel(dbx).rename(columns={'Unnamed: 0': 'ID'})
db = db[pd.to_numeric(db.ID, errors='coerce').notna()]
db['ID'] = db.ID.astype(int)
db = db.set_index('ID')
col = {c.split('\n')[0].split(' (')[0].strip(): c for c in db.columns}
AREA_ATENEO = db['Area Assegnata Ateneo']

ids = sorted(R.index)


def categoria(e):
    if e.startswith('Spostamento piccolo'):
        return 'Spostamento <=50 m'
    if e.startswith('Spostamento possibile'):
        return 'Spostamento 50-200 m'
    if e.startswith('Area DB sotto'):
        return 'Area DB sotto soglia'
    return 'Probabile errore database'


def tipo_errore(i):
    """errore nel valore dell'area o nella posizione, dal confronto con la larghezza dell'alveo"""
    adb, acalc = NC.loc[i, 'Area database (km²)'], NC.loc[i, 'Area calcolata 5 m (km²)']
    aw = NC.loc[i, 'Area stimata da larghezza alveo (km²)']
    rete = str(NC.loc[i, 'Tipologia reticolo (DB)'])
    if 'artificiale' in rete:
        return ('Bacino non ricostruibile dal DEM',
                'Il database indica un reticolo artificiale di scolo: il bacino e\' definito dalla rete di '
                'bonifica e non dalla topografia, quindi il DEM non puo\' ricostruirlo.')
    orig = str(NC.loc[i, 'Probabile origine errore'])
    if orig.startswith('Coordinate'):
        return ('Coordinate/posizione del ponte',
                'Le coordinate del ponte nel database sono incongruenti (gia\' emerso nel controllo coordinate: '
                + str(NC.loc[i, 'Segnalazioni']).split(';')[0] + '): il bacino calcolato non e\' quello del ponte.')
    if 'ripetuta' in str(NC.loc[i, 'Segnalazioni']) or (adb <= 0.1 and pd.isna(aw)):
        return ('Valore di area del database',
                f'L\'area DB di {adb:g} km² e\' un valore ripetuto su piu\' ponti (probabile segnaposto) e la '
                'larghezza dell\'alveo non e\' riportata.')
    if pd.isna(aw):
        return ('Da verificare', 'Larghezza dell\'alveo non riportata nel database: non e\' possibile capire quale dato sia errato.')
    ldb, lc = abs(np.log10(aw / adb)), abs(np.log10(aw / acalc))
    testo_w = (f'La larghezza dell\'alveo inciso riportata nel database ({NC.loc[i, "Larghezza alveo inciso (m, DB)"]:g} m) '
               f'corrisponde a un bacino di circa {aw:.3g} km²')
    if ldb <= 0.5 and ldb < lc - 0.3:
        return ('Coordinate/posizione del ponte',
                testo_w + f', coerente con l\'area DB ({adb:g} km²) e non con quella calcolata ({acalc:.3g} km²): '
                'il ponte attraversa un corso delle dimensioni indicate dal database, ma vicino alle coordinate il '
                'DEM non ne trova uno. L\'errore e\' probabilmente nella posizione (coordinate DB/KML), non nell\'area.')
    if lc <= 0.5 and lc < ldb - 0.3:
        return ('Valore di area del database',
                testo_w + f', coerente con l\'area calcolata ({acalc:.3g} km²) e non con quella DB ({adb:g} km²): '
                'l\'errore e\' probabilmente nel valore di area del database.')
    return ('Da verificare',
            testo_w + f': non conferma chiaramente ne\' l\'area DB ({adb:g} km²) ne\' quella calcolata ({acalc:.3g} km²).')


righe, nuove, linee = [], [], []
for i in ids:
    r = R.loc[i]
    cat = categoria(r['Esito della ricerca'])
    s = S.geometry[i]
    rec = {'ID': i, 'Area Assegnata Ateneo': AREA_ATENEO.get(i), 'Strada': NC.loc[i, 'Strada'],
           'Progressiva': NC.loc[i, 'Progressiva'], 'Nome opera': NC.loc[i, 'Nome opera'],
           'Tipologia reticolo (DB)': NC.loc[i, 'Tipologia reticolo (DB)'],
           'Larghezza alveo inciso (m, DB)': NC.loc[i, 'Larghezza alveo inciso (m, DB)'],
           'Area database (km²)': r['Area database (km²)'], 'Area calcolata (km²)': r['Area calcolata (km²)'],
           'Differenza %': r['Differenza %'],
           'Area stimata da larghezza alveo (km²)': NC.loc[i, 'Area stimata da larghezza alveo (km²)'],
           'Categoria': cat}
    nuova = s
    a_new, spost = r['Area calcolata (km²)'], 0.0
    if cat != 'Probabile errore database' and i in C.index:
        nuova = C.geometry[i]
        a_new = C.area_km2[i]
        spost = C.dist_m[i]
        linee.append({'ID': i, 'categoria': cat, 'spost_m': round(float(s.distance(nuova)), 1),
                      'geometry': LineString([(s.x, s.y), (nuova.x, nuova.y)])})
    rec['Area alla nuova posizione (km²)'] = a_new
    rec['Differenza alla nuova posizione %'] = a_new / r['Area database (km²)'] - 1
    rec['Distanza della nuova sezione dal ponte (m)'] = spost
    rec['Corso della nuova sezione'] = r.get('Sezione compatibile più vicina: corso') if cat != 'Probabile errore database' else ''
    d = r.get('Sezione compatibile più vicina: distanza dal ponte (m)')
    if cat == 'Probabile errore database':
        te, motivo = tipo_errore(i)
        vic = ('nessuna cella del reticolo con area compatibile (±15%) entro circa 750 m dal ponte'
               if pd.isna(d) else f'la sezione con area compatibile piu\' vicina e\' a {d:.0f} m dal ponte '
               f'({r["Sezione compatibile più vicina: corso"]})')
        if 'Idrovora' in str(db.loc[i, 'DENOMINAZIONE GOOGLE EARTH']):
            motivo += (' Il toponimo (Via Idrovora) indica un\'area di bonifica con idrovora: li\' il bacino e\' '
                       'regolato dalla rete di scolo e il DEM non lo ricostruisce.')
        rec['Tipo di errore nel database'] = te
        rec['_breve'] = ('nessuna sez. compatibile entro 750 m' if pd.isna(d) else f'sez. compatibile a {d:.0f} m') + \
            f'; errore DB: {te}'
        rec['Giustificazione'] = (f'Area DB {r["Area database (km²)"]:g} km², calcolata {r["Area calcolata (km²)"]:.3g} km² '
                                  f'({r["Differenza %"]:+.0%}); {vic}. Spostare la sezione fin li\' significherebbe '
                                  'calcolare il bacino di un altro punto o di un altro corso, non quello del ponte. '
                                  + motivo)
    else:
        rec['Tipo di errore nel database'] = ''
        rec['Giustificazione'] = ''
        rec['_breve'] = ''
    righe.append(rec)
    nuove.append({'ID': i, 'geometry': nuova})

T = pd.DataFrame(righe)

# ---- shapefile ----
def attr(df):
    return pd.DataFrame({
        'ID': df.ID, 'area_aten': df['Area Assegnata Ateneo'].fillna(''), 'strada': df.Strada.fillna('').astype(str),
        'progr': df.Progressiva.fillna('').astype(str), 'nome': df['Nome opera'].fillna('').astype(str).str[:80],
        'A_db_km2': df['Area database (km²)'], 'A_calc_km2': df['Area calcolata (km²)'],
        'diff_pct': (df['Differenza %'] * 100).round(1), 'categoria': df.Categoria})


A = attr(T)
gpd.GeoDataFrame(A, geometry=list(S.geometry.loc[ids]), crs=32633).to_file(
    os.path.join(out, 'ponti_nc_posizione_attuale.shp'), encoding='utf-8')
B = attr(T)
B['A_new_km2'] = T['Area alla nuova posizione (km²)'].round(3)
B['diff_new'] = (T['Differenza alla nuova posizione %'] * 100).round(1)
B['spost_m'] = T['Distanza della nuova sezione dal ponte (m)'].round(1)
B['corso'] = T['Corso della nuova sezione'].fillna('')
B['tipo_err'] = T['Tipo di errore nel database'].fillna('')
B['giustif'] = T['_breve'].str[:254]  # testo completo nell'Excel
gpd.GeoDataFrame(B, geometry=[n['geometry'] for n in nuove], crs=32633).to_file(
    os.path.join(out, 'ponti_nc_posizione_nuova.shp'), encoding='utf-8')
gpd.GeoDataFrame(linee, crs=32633).to_file(os.path.join(out, 'spostamenti_sezione.shp'), encoding='utf-8')

# ---- stili QGIS (.qml con lo stesso nome dello shapefile: QGIS lo applica da solo) ----
def simbolo(nome, colore, forma='circle', dim='3.6'):
    return f'''<symbol name="{nome}" type="marker" alpha="1"><layer class="SimpleMarker" pass="0" enabled="1" locked="0">
<Option type="Map"><Option name="color" type="QString" value="{colore},255"/><Option name="name" type="QString" value="{forma}"/>
<Option name="size" type="QString" value="{dim}"/><Option name="outline_color" type="QString" value="35,35,35,255"/>
<Option name="outline_width" type="QString" value="0.3"/></Option></layer></symbol>'''


def qml_categorie(cat_colori, forma='circle', dim='3.6', etichetta=True):
    cats = ''.join(f'<category render="true" symbol="{k}" value="{v}" label="{v}"/>' for k, (v, _) in enumerate(cat_colori))
    syms = ''.join(simbolo(str(k), c, forma, dim) for k, (_, c) in enumerate(cat_colori))
    lab = '''<labeling type="simple"><settings calloutType="simple"><text-style fieldName="ID" fontSize="9" textColor="0,0,0,255">
<text-buffer bufferDraw="1" bufferSize="0.8" bufferColor="255,255,255,255"/></text-style>
<placement placement="0" dist="1.2"/><rendering scaleVisibility="0"/></settings></labeling>''' if etichetta else ''
    return f'''<!DOCTYPE qgis PUBLIC 'http://mrcc.com/qgis.dtd' 'SYSTEM'>
<qgis version="3.28" styleCategories="Symbology|Labeling" labelsEnabled="{1 if etichetta else 0}">
<renderer-v2 type="categorizedSymbol" attr="categoria" enableorderby="0" forceraster="0" symbollevels="0">
<categories>{cats}</categories><symbols>{syms}</symbols></renderer-v2>{lab}</qgis>'''


COL = [('Probabile errore database', '215,25,28'), ('Spostamento <=50 m', '26,150,65'),
       ('Spostamento 50-200 m', '253,174,97'), ('Area DB sotto soglia', '43,131,186')]
open(os.path.join(out, 'ponti_nc_posizione_nuova.qml'), 'w').write(qml_categorie(COL))
open(os.path.join(out, 'ponti_nc_posizione_attuale.qml'), 'w').write(
    qml_categorie([(v, '150,150,150') for v, _ in COL], forma='square', dim='3.0', etichetta=False))
open(os.path.join(out, 'spostamenti_sezione.qml'), 'w').write('''<!DOCTYPE qgis PUBLIC 'http://mrcc.com/qgis.dtd' 'SYSTEM'>
<qgis version="3.28" styleCategories="Symbology"><renderer-v2 type="singleSymbol"><symbols>
<symbol name="0" type="line" alpha="1"><layer class="SimpleLine" pass="0" enabled="1" locked="0"><Option type="Map">
<Option name="line_color" type="QString" value="60,60,60,255"/><Option name="line_width" type="QString" value="0.5"/>
<Option name="line_style" type="QString" value="dash"/></Option></layer></symbol></symbols></renderer-v2></qgis>''')

# ---- Excel ----
MOTIVO = [
    'Perche\' non conviene spostare la sezione per i ponti classificati "Probabile errore database"',
    '',
    '1. Per questi ponti non esiste, entro 200 m dal ponte (tracciato KML e punti di inizio/fine del database), '
    'alcuna sezione del reticolo il cui bacino abbia un\'area compatibile (±15%) con quella del database. '
    'Per 10 di essi non ne esiste nessuna entro circa 750 m.',
    '2. Ottenere la compatibilita\' richiederebbe di spostare la sezione di centinaia di metri o su un altro corso '
    'd\'acqua: si calcolerebbe il bacino di un altro punto, non quello del ponte. Sarebbe una taratura sul dato '
    'che si vuole verificare, e il risultato non avrebbe piu\' valore di controllo.',
    '3. L\'area calcolata alla sezione attuale deriva dallo stesso metodo che rende compatibili 168 ponti su 207 '
    '(DEM 5 m, reticolo di riferimento, intersezione con il tracciato del ponte, soglia scelta con validazione '
    'incrociata): non c\'e\' motivo di ritenerla meno affidabile per questi ponti.',
    '4. Un controllo indipendente, la larghezza dell\'alveo inciso riportata nel database (relazione larghezza-area '
    'ricavata dai ponti compatibili), indica per ciascun ponte quale dato e\' errato (colonna "Tipo di errore nel '
    'database"): il valore di area, oppure le coordinate/posizione del ponte (il corso d\'acqua con l\'area del '
    'database esiste ma non passa vicino alle coordinate).',
    '5. Si raccomanda di mantenere l\'area calcolata con la segnalazione e di correggere il database (area o '
    'coordinate) dopo verifica su ortofoto o sopralluogo.',
    '',
    'Gli spostamenti proposti per le altre categorie (<=50 m e 50-200 m sullo stesso corso) correggono invece la '
    'posizione della sezione di pochi metri sul corso attraversato dal ponte; quelli su un altro corso o su un '
    'affluente vanno confermati sulla cartografia prima di adottarli.',
    '',
    'Area di appartenenza: colonna "Area Assegnata Ateneo" del database (Centrale, Nord-Est, Sud-Est).',
]
E = T[T.Categoria == 'Probabile errore database']
xo = os.path.join(out, 'Giustificazione_non_compatibili.xlsx')
with pd.ExcelWriter(xo, engine='openpyxl') as w:
    pd.DataFrame({'Motivazione': MOTIVO}).to_excel(w, sheet_name='Motivazione', index=False)
    w.sheets['Motivazione'].column_dimensions['A'].width = 160
    cols_e = ['ID', 'Area Assegnata Ateneo', 'Strada', 'Progressiva', 'Nome opera', 'Tipologia reticolo (DB)',
              'Larghezza alveo inciso (m, DB)', 'Area database (km²)', 'Area calcolata (km²)', 'Differenza %',
              'Area stimata da larghezza alveo (km²)', 'Tipo di errore nel database', 'Giustificazione']
    E[cols_e].to_excel(w, sheet_name='Errori nel database (19)', index=False)
    cols_t = ['ID', 'Area Assegnata Ateneo', 'Strada', 'Progressiva', 'Nome opera', 'Area database (km²)',
              'Area calcolata (km²)', 'Differenza %', 'Categoria', 'Distanza della nuova sezione dal ponte (m)',
              'Corso della nuova sezione', 'Area alla nuova posizione (km²)', 'Differenza alla nuova posizione %',
              'Tipo di errore nel database']
    T[cols_t].to_excel(w, sheet_name='Tutti i non compatibili (39)', index=False)
    P = pd.crosstab(T['Area Assegnata Ateneo'], T.Categoria, margins=True, margins_name='Totale')
    P.to_excel(w, sheet_name='Per area Ateneo')
    for nome, df in (('Errori nel database (19)', E[cols_e]), ('Tutti i non compatibili (39)', T[cols_t])):
        ws = w.sheets[nome]
        for k, c in enumerate(df.columns, 1):
            L = ws.cell(row=1, column=k).column_letter
            ws.column_dimensions[L].width = 100 if c == 'Giustificazione' else min(max(len(c), 10) + 2, 40)
            if c.endswith('%'):
                for row in ws.iter_rows(min_row=2, min_col=k, max_col=k):
                    row[0].number_format = '0.0%'
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = ws.dimensions
    from openpyxl.styles import Alignment
    ws = w.sheets['Errori nel database (19)']
    jg = cols_e.index('Giustificazione') + 1
    for row in ws.iter_rows(min_row=2, min_col=jg, max_col=jg):
        row[0].alignment = Alignment(wrap_text=True, vertical='top')
    w.sheets['Per area Ateneo'].column_dimensions['A'].width = 24

print(T.Categoria.value_counts().to_string())
print(E[['ID', 'Area Assegnata Ateneo', 'Tipo di errore nel database']].to_string(index=False))
print(P.to_string())
