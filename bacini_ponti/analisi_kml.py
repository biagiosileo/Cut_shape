"""
Controlli sui ponti fuori tolleranza con le informazioni del KML (localizzazione, tracciati, confini dei
comuni e delle aree dell'Ateneo) e dell'elenco complessivo dei ponti (comune, intercetta asta, note),
piu' la coerenza fisica tra luce del ponte / larghezza dell'alveo e area del bacino.

Controlli per ponte:
  - comune che contiene i punti inizio e fine del database, il punto del KML e la sezione di calcolo,
    confrontato con il comune dell'elenco;
  - area dell'Ateneo della cartella del KML e del confine nel KML, confrontata con quella del database;
  - lunghezza del tracciato del KML confrontata con la lunghezza del ponte nel database;
  - "Intercetta asta fluviale/canale?" e note dell'elenco;
  - luce e larghezza dell'alveo inciso: sui ponti compatibili si stima la relazione
    log(larghezza) = a + b log(area) e si dice se l'area del database o quella calcolata e' piu'
    coerente con la larghezza del ponte (residuo in unita' di deviazione standard).

Uso: python analisi_kml.py <doc.kml> <Elenco_ponti.xlsx> <Database.xlsx> <controllo_coordinate.gpkg>
        <ponti.csv> <parametri.csv> <revisione.csv> <out_dir>
"""
import os
import re
import sys
import xml.etree.ElementTree as ET

import geopandas as gpd
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from shapely.geometry import LineString, Point, Polygon
from shapely.ops import unary_union

kml, elenco, dbx, gpkg, pcsv, par, rev, out = sys.argv[1:9]
os.makedirs(out, exist_ok=True)
NS = '{http://www.opengis.net/kml/2.2}'
FUORI = [6, 13, 19, 21, 44, 58, 64, 72, 83, 85, 90, 102, 103, 104, 105, 106, 107, 108, 109, 111, 113, 116, 117,
         118, 124, 129, 143, 151, 154, 155, 164, 168, 174, 177, 190, 196, 198, 200, 202]
AREE_ATENEO = {'area sud-est': 'Sud-Est', 'area nord-est': 'Nord-Est', 'area centrale': 'Centrale'}


def xy(s):
    return [tuple(map(float, p.split(',')[:2])) for p in s.split()]


# --- KML: poligoni (comuni e aree dell'Ateneo), punti e tracciati con la cartella ---
poli, pm = [], []
for f in ET.parse(kml).getroot().iter(NS + 'Placemark'):
    nome = (f.findtext(NS + 'name') or '').strip()
    for g in f.iter(NS + 'Polygon'):
        poli.append({'nome': nome, 'geometry': Polygon(xy(g.find(f'.//{NS}outerBoundaryIs//{NS}coordinates').text))})
Pol = gpd.GeoDataFrame(poli, crs=4326).to_crs(32633)
Pol['geometry'] = Pol.buffer(0)
Pol['chiave'] = Pol.nome.str.replace(r'\s+\d+$', '', regex=True).str.strip().str.lower()
aree = Pol[Pol.chiave.isin(AREE_ATENEO)].dissolve('chiave').geometry
comuni = Pol[~Pol.chiave.isin(AREE_ATENEO)].dissolve('chiave').geometry


def comune_di(p):
    if p is None or p.is_empty:
        return ''
    c = comuni[comuni.contains(p)]
    return c.index[0].title() if len(c) else 'fuori dai comuni del KML'


def area_di(p):
    c = aree[aree.contains(p)]
    return ', '.join(AREE_ATENEO[k] for k in c.index) if len(c) else 'fuori'


# --- elenco complessivo: comune, intercetta asta, note (abbinamento strada + progressiva) ---
El = pd.read_excel(elenco, header=0)
El.columns = [str(v).strip() for v in El.iloc[0]]
El = El.iloc[2:]
El = El.loc[:, ~El.columns.duplicated()]
chiave = lambda s, k: str(s).strip().upper().replace(' ', '') + '|' + re.sub(r'\s', '', str(k))
El['k'] = [chiave(s, k) for s, k in zip(El['NUOVA NUMERAZIONE'], El['Progressiva km'])]
ci = [c for c in El.columns if c.lower().startswith('intercetta')][0]
El = El.drop_duplicates('k').set_index('k')

db = pd.read_excel(dbx).rename(columns={'Unnamed: 0': 'ID'})
db = db[pd.to_numeric(db.ID, errors='coerce').notna()]
db['ID'] = db.ID.astype(int)
db = db.set_index('ID')
num = lambda c: pd.to_numeric(db[c].astype(str).str.replace(',', '.').str.extract(r'([\d.]+)')[0], errors='coerce')
adb = num('SUPERFICIE BACINO (km2)')
luce, alveo, lung = num('LUCE TOT (SPALLA/SPALLA - m)'), num('LARGHEZZA ALVEO INCISO (m)'), num(
    'LUNGHEZZA COMPLESSIVA (INCLUSE LE SPALLE - m)')

ini = gpd.read_file(gpkg, layer='db_inizio').set_index('ID').to_crs(32633).geometry
fin = gpd.read_file(gpkg, layer='db_fine').set_index('ID').to_crs(32633).geometry
kp = gpd.read_file(gpkg, layer='kml_abbinato').set_index('ID').to_crs(32633)
kt = gpd.read_file(gpkg, layer='kml_tracciati').to_crs(32633).set_index('nome')
P = pd.read_csv(pcsv).set_index('id')
sez = gpd.GeoSeries(gpd.points_from_xy(P.X, P.Y), index=P.index, crs=4326).to_crs(32633)
U = pd.read_csv(par)
U['ID'] = U.ID.astype(int)
acalc = U.set_index('ID').A_km2
Rv = pd.read_csv(rev).set_index('ID')
cart = {}
for f in ET.parse(kml).getroot().iter(NS + 'Folder'):
    n = (f.findtext(NS + 'name') or '').lower()
    for p in f.iter(NS + 'Placemark'):
        cart[(p.findtext(NS + 'name') or '').strip()] = n

# --- relazione larghezza - area sui ponti compatibili ---
comp = [i for i in adb.index if i in acalc.index and pd.notna(adb[i]) and abs(acalc[i] / adb[i] - 1) <= 0.15]
larg = alveo.where(alveo > 0, luce)
ok = [i for i in comp if pd.notna(larg.get(i)) and larg[i] > 0]
b, a = np.polyfit(np.log10(adb[ok].astype(float)), np.log10(larg[ok].astype(float)), 1)
sd = float(np.std(np.log10(larg[ok].astype(float)) - (a + b * np.log10(adb[ok].astype(float)))))


def z(area, i):
    return (np.log10(larg[i]) - (a + b * np.log10(area))) / sd


righe = []
for i in FUORI:
    k = chiave(db.loc[i, 'NUOVA NUMERAZIONE'], db.loc[i, 'Progressiva km'])
    e = El.loc[k] if k in El.index else None
    knome = kp.loc[i, 'KML_nome']
    tr = kt.geometry.get(knome)
    c_el = str(e['Comune']).strip() if e is not None else ''
    cc = {'inizio DB': comune_di(ini.get(i)), 'fine DB': comune_di(fin.get(i)), 'punto KML': comune_di(kp.geometry[i]),
          'sezione': comune_di(sez[i])}
    diversi = [n for n, v in cc.items() if c_el and v.lower() != c_el.lower()]
    a_cart = next((v for kk, v in AREE_ATENEO.items() if kk in cart.get(knome, '')), '')
    zdb = z(adb[i], i) if pd.notna(larg.get(i)) and pd.notna(adb[i]) else np.nan
    zc = z(acalc[i], i) if pd.notna(larg.get(i)) else np.nan
    if np.isnan(zdb):
        coer = 'larghezza non disponibile'
    elif abs(zdb) <= 1.5 and abs(zc) > 1.5:
        coer = "piu' coerente l'area del database"
    elif abs(zc) <= 1.5 and abs(zdb) > 1.5:
        coer = "piu' coerente l'area calcolata"
    elif abs(zc) <= 1.5 and abs(zdb) <= 1.5:
        coer = 'entrambe coerenti'
    else:
        coer = 'nessuna delle due coerente'
    righe.append({
        'ID': i, 'Nome opera': db.loc[i, 'Nome opera'], 'SP': db.loc[i, 'NUOVA NUMERAZIONE'], 'km': db.loc[i, 'Progressiva km'],
        'A_db_km2': adb[i], 'A_calcolata_km2': acalc[i], 'Gruppo revisione': Rv.Gruppo.get(i, 'DEM'),
        'Comune elenco': c_el, 'Comune inizio DB': cc['inizio DB'], 'Comune fine DB': cc['fine DB'],
        'Comune punto KML': cc['punto KML'], 'Comune sezione': cc['sezione'],
        'Comune diverso da elenco': ', '.join(diversi),
        'Area Ateneo DB': db.loc[i, 'Area Assegnata Ateneo'], 'Area Ateneo cartella KML': a_cart,
        'Area Ateneo (confine KML) sezione': area_di(sez[i]),
        'Lunghezza ponte DB (m)': lung.get(i), 'Lunghezza tracciato KML (m)': round(tr.length) if tr is not None else None,
        'Luce tot DB (m)': luce.get(i), 'Larghezza alveo inciso DB (m)': alveo.get(i),
        'Tipologia reticolo DB': db.loc[i, 'TIPOLOGIA RETICOLO (principale / secondario)'],
        'Intercetta asta (elenco)': e[ci] if e is not None else '', 'Note (elenco)': e['Note'] if e is not None and 'Note' in e else '',
        'Scarto larghezza con area DB (sd)': round(zdb, 1) if not np.isnan(zdb) else None,
        'Scarto larghezza con area calcolata (sd)': round(zc, 1) if not np.isnan(zc) else None,
        'Coerenza larghezza-area': coer})
T = pd.DataFrame(righe)
xl = os.path.join(out, 'Analisi_KML_non_compatibili.xlsx')
nota = pd.DataFrame({'Relazione larghezza alveo (o luce) - area sui ponti compatibili': [
    f'log10(larghezza) = {a:.3f} + {b:.3f} log10(area), ponti usati {len(ok)}, deviazione standard {sd:.3f}',
    "Scarto (sd) = quanto la larghezza del ponte si discosta dalla larghezza attesa per quell'area; entro +-1,5 = coerente.",
    "Comuni e aree dell'Ateneo dai poligoni del KML; comune di riferimento dall'elenco complessivo dei ponti."]})
with pd.ExcelWriter(xl, engine='openpyxl') as w:
    T.to_excel(w, sheet_name='Controlli 39 ponti', index=False)
    nota.to_excel(w, sheet_name='Metodo', index=False)
wb = load_workbook(xl)
for ws in wb.worksheets:
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDDDDD')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.freeze_panes = 'C2'
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 16
wb['Metodo'].column_dimensions['A'].width = 130
wb.save(xl)
pd.set_option('display.width', 250)
print(f'relazione: a={a:.3f} b={b:.3f} sd={sd:.3f} n={len(ok)}')
print(T.drop(columns=['Nome opera', 'Note (elenco)']).to_string())
print(T[['ID', 'Note (elenco)', 'Intercetta asta (elenco)']].to_string(max_colwidth=120))
