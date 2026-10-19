"""
Risultati della correzione delle porzioni (cartella separata, i risultati precedenti non cambiano):
  - parametri_ponti_porzioni_corrette.csv  parametri del codice dell'utente per i ponti ricalcolati
  - bacini_porzioni_corrette.shp           bacini ricalcolati (campo ID = ID del database)
  - Correzione_porzioni.xlsx               confronto prima/dopo con l'area del database

Uso: python correzione_porzioni_risultati.py <cartella porzioni nuove> <ponti_porzioni_corrette.csv>
        <ponti.csv originale> <parametri prima (csv del codice)> <Database.xlsx> <out_dir>
"""
import glob
import json
import os
import sys

import geopandas as gpd
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from shapely.geometry import shape

cart, csv_nuovo, csv_vecchio, par_prima, dbx, out = sys.argv[1:7]
N = pd.read_csv(csv_nuovo).set_index('id')
V = pd.read_csv(csv_vecchio).set_index('id')
ids = sorted(N.index[N.porzione != V.porzione.reindex(N.index)])

par = pd.concat([pd.read_csv(f) for f in sorted(glob.glob(os.path.join(cart, '*_parametri.csv')))])
par['ID'] = par.ID.astype(int)
par = par[par.ID.isin(ids)].sort_values('ID')
par.to_csv(os.path.join(out, 'parametri_ponti_porzioni_corrette.csv'), index=False)

feats = []
for f in sorted(glob.glob(os.path.join(cart, '*_bacini_afferenti.geojson'))):
    feats += json.load(open(f))['features']
g = gpd.GeoDataFrame({'ID': [int(ft['properties']['ID']) for ft in feats]},
                     geometry=[shape(ft['geometry']) for ft in feats], crs=32633)
g = g[g.ID.isin(ids)].sort_values('ID')
g['A_km2'] = (g.area / 1e6).round(3)
g['porzione'] = [N.loc[i, 'porzione'] for i in g.ID]
g.to_file(os.path.join(out, 'bacini_porzioni_corrette.shp'))

prima = pd.read_csv(par_prima)
prima['ID'] = prima.ID.astype(int)
prima = prima.set_index('ID')
db = pd.read_excel(dbx).rename(columns={'Unnamed: 0': 'ID'})
db = db[pd.to_numeric(db.ID, errors='coerce').notna()]
db['ID'] = db.ID.astype(int)
db = db.set_index('ID')
dopo = par.set_index('ID')


def classe(a, adb):
    if pd.isna(adb):
        return 'senza area DB'
    d = abs(a / adb - 1)
    return 'compatibile' if d <= .15 else ('quasi compatibile' if d <= .25 else 'non compatibile')


NOTE = {
    6: "Porzione troppo stretta (4,6 km): il bacino dell'Agri era tagliato. Con la porzione corretta il bacino si "
       "chiude contro il limite ovest del DEM unico (circa E 576 km): l'alto Agri non e' coperto dai DEM disponibili, "
       "l'area e' per difetto. L'area del database e' plausibile.",
    13: "Il bacino toccava il bordo solo sullo spartiacque: area quasi invariata. Resta la differenza con il database.",
    154: "Il bacino toccava il bordo solo sullo spartiacque: area invariata. Resta la differenza con il database.",
    155: "Il bacino toccava il bordo solo sullo spartiacque: area invariata. Resta la differenza con il database.",
    15: "Il bacino toccava il bordo solo sullo spartiacque: area quasi invariata. Area non presente nel database.",
    109: "Bacino troncato dalla porzione: con la porzione corretta l'area cresce e si allontana ancora di piu' dal "
         "valore del database (0,31 km2).",
    29: "Bacino leggermente troncato a sud: con la porzione corretta l'area si avvicina al database.",
    73: "Bacino leggermente troncato a sud: con la porzione corretta l'area si avvicina al database.",
    79: "Bacino leggermente troncato a sud: con la porzione corretta l'area si avvicina al database.",
}
righe = []
for i in ids:
    adb = pd.to_numeric(db.loc[i, 'SUPERFICIE BACINO (km2)'], errors='coerce')
    nome = str(db.loc[i, 'Nome opera']).strip()
    righe.append([i, nome if nome not in ('-', 'nan') else '', V.loc[i, 'porzione'], N.loc[i, 'porzione'],
                  None if pd.isna(adb) else float(adb), float(prima.loc[i, 'A_km2']), float(dopo.loc[i, 'A_km2']),
                  None, None, classe(prima.loc[i, 'A_km2'], adb), classe(dopo.loc[i, 'A_km2'], adb), NOTE.get(i, '')])
cols = ['ID', 'Nome opera', 'Porzione prima', 'Porzione nuova', 'Area DB (km²)', 'Area prima (km²)',
        'Area dopo (km²)', 'Diff. prima vs DB', 'Diff. dopo vs DB', 'Classe prima', 'Classe dopo', 'Nota']
T = pd.DataFrame(righe, columns=cols)
por = N.loc[ids].reset_index().groupby('porzione').agg(
    ponti=('id', lambda s: ', '.join(map(str, s))), xmin=('xmin', 'first'), ymin=('ymin', 'first'),
    xmax=('xmax', 'first'), ymax=('ymax', 'first')).reset_index()
por['celle (milioni)'] = ((por.xmax - por.xmin) * (por.ymax - por.ymin) / 25 / 1e6).round(1)

xl = os.path.join(out, 'Correzione_porzioni.xlsx')
with pd.ExcelWriter(xl, engine='openpyxl') as w:
    T.to_excel(w, sheet_name='Confronto aree', index=False)
    por.to_excel(w, sheet_name='Porzioni nuove', index=False)
wb = load_workbook(xl)
ws = wb['Confronto aree']
for r in range(2, ws.max_row + 1):
    ws[f'H{r}'] = f'=IF(E{r}="","",F{r}/E{r}-1)'
    ws[f'I{r}'] = f'=IF(E{r}="","",G{r}/E{r}-1)'
    for c in 'HI':
        ws[f'{c}{r}'].number_format = '0.0%'
    for c in 'EFG':
        ws[f'{c}{r}'].number_format = '0.000'
    ws[f'L{r}'].alignment = Alignment(wrap_text=True, vertical='top')
for s in wb.worksheets:
    for c in s[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDDDDD')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    s.freeze_panes = 'A2'
for c, wdt in zip('ABCDEFGHIJKL', [6, 22, 14, 14, 11, 11, 11, 11, 11, 16, 16, 80]):
    ws.column_dimensions[c].width = wdt
wb['Porzioni nuove'].column_dimensions['B'].width = 14
wb.save(xl)
print(T.drop(columns='Nota').to_string())
