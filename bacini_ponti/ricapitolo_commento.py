"""
Commento al ricapitolo (Word): stato dei ponti, cosa e' stato corretto, cosa resta e come procedere,
con i criteri per spostare la sezione di un ponte. Times New Roman 11 pt, testo giustificato.

Uso: python ricapitolo_commento.py <Ricapitolo_confronto_aree.xlsx> <test_v2.csv> <out.docx>
"""
import sys

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

xl, test, out = sys.argv[1:4]
NC = pd.read_excel(xl, sheet_name='Non compatibili - azioni')
RC = pd.read_excel(xl, sheet_name='Ricapitolo correzioni')
T = pd.read_csv(test).set_index('ID')

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(8)


def par(t, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(t)
    r.bold = bold
    return p


def titolo(t):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold = True
    r.font.size = Pt(11)


def tabella(df, cols):
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = 'Table Grid'
    for j, c in enumerate(cols):
        t.rows[0].cells[j].text = c
        t.rows[0].cells[j].paragraphs[0].runs[0].bold = True
    for _, r in df.iterrows():
        cells = t.add_row().cells
        for j, c in enumerate(cols):
            v = r[c]
            cells[j].text = (f'{v:.3f}' if isinstance(v, float) and abs(v) < 10 else
                             f'{v:,.1f}'.replace(',', 'X').replace('.', ',').replace('X', '.') if isinstance(v, float) else str(v))
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.name = 'Times New Roman'
    doc.add_paragraph()


a6, a105 = T.loc[6, 'A_finale_km2'], T.loc[105, 'A_finale_km2']
titolo('Ricapitolo della verifica delle aree dei bacini dei 212 ponti')
titolo('1. Stato attuale')
par('Il confronto e\' tra l\'area calcolata dal codice (pysheds_bridges_5m.py sul DEM unico a 5 m) e la SUPERFICIE BACINO '
    'del database. Classi: compatibile se la differenza e\' entro il ±15%; quasi compatibile se e\' entro il ±25%, oppure, per '
    'i bacini piccoli, se la differenza assoluta e\' entro 0,1 km² (la soglia del reticolo) e quella percentuale entro il 50%; '
    'non compatibile altrimenti.')
par('Dalla prima esecuzione al run definitivo i ponti compatibili sono passati da 168 a 175, i quasi compatibili da 11 a 12 e '
    'i non compatibili da 28 a 20 (5 ponti non hanno l\'area nel database). Il run definitivo e\' stato verificato: 212 righe, '
    'ID 1-212 in ordine, coordinate e parametri identici a quelli attesi.')
titolo('2. Cosa e\' stato corretto')
par(f'Rispetto al ponti.csv originale sono stati modificati {len(RC)} ponti, con tre tipi di correzione: porzioni di DEM '
    'allargate dove il bacino era troncato dal bordo della porzione; sezione spostata sul punto in cui il tracciato del ponte '
    'nel KML attraversa il reticolo; soglia di aggancio ridotta per i fossi piu\' piccoli della soglia del reticolo. Con queste '
    'correzioni sono diventati compatibili 64, 83, 90, 143, 168, 177 e 202; 103, 116, 129 e 151 sono passati da non compatibili '
    'a quasi compatibili; 29, 73 e 79 si sono avvicinati al database. L\'elenco completo e\' nel foglio Ricapitolo correzioni.')
par('I quasi compatibili (44, 72, 85, 102, 103, 104, 106, 116, 124, 129, 151, 200) restano tali: le differenze sono di poche '
    'centinaia di metri quadri o di pochi ettari, dentro l\'incertezza del DEM a 5 m e del valore del database, e spostarli non '
    'avrebbe un fondamento sufficiente.')
titolo('3. I 20 non compatibili')
tabella(NC, ['ID', 'Gruppo', 'Area DB (km²)', 'Area calcolata (km²)', 'Cosa fare'])
par(f'DEM esteso (ponte 6). Il bacino dell\'Agri usciva dal DEM unico a ovest di E 576 km. Nessuno dei DEM forniti copre '
    f'l\'alto Agri; per la prova e\' stato usato il DEM pubblico Copernicus GLO-30 (30 m) per riempire 575 km² di celle senza dato, '
    f'riportato sulla griglia a 5 m e corretto della differenza mediana con il DEM a 5 m (1,4 m), con il reticolo di '
    f'riferimento inciso come nel resto del DEM. Con il DEM unico v2 il bacino del ponte 6 risulta di {a6:.1f} km² '
    f'(database 1609,8 km², {100 * (a6 / 1609.77 - 1):+.1f}%). Il Copernicus e\' un modello di superficie (include la '
    f'vegetazione) a 30 m: e\' adeguato per l\'area di un bacino cosi\' grande, meno per i parametri di dettaglio; se si trova un '
    f'DTM a 5 m dell\'alto Agri conviene sostituirlo, il vettoriale di ritaglio e\' gia\' disponibile.')
par(f'Tombino (ponte 105). Il rilevato stradale fa da diga nel DEM: l\'acqua del fosso non passa sotto la strada e scorre '
    f'lungo il piede del rilevato. Nel DEM unico v2 e\' inciso un taglio di 60 m perpendicolare alla strada in corrispondenza '
    f'del ponte; con la sezione sul taglio il bacino vale {a105:.3f} km² (database 0,12 km²).')
par('Verifica a vista (107, 111). Per il 107 il ponte e\' sulla confluenza di due rami: va visto quale passa sotto il ponte. '
    'Per il 111 il tracciato KML corre parallelo al torrente a 16 m; il fosso sotto il tracciato ha solo 0,15 km², mentre il '
    'torrente (15,8 km²) e la larghezza dell\'alveo del database (32 m) sono coerenti con i 14,5 km² del database: se '
    'l\'ortofoto conferma che il ponte scavalca il torrente, la sezione va spostata sul torrente.')
par('Bonifica e canali (13, 154, 155, 164). In pianura il deflusso e\' deciso dai canali e dai manufatti, che il DEM a 5 m '
    'non vede. Servono il reticolo dei canali del Consorzio di Bonifica o l\'idrografia della CTR, da incidere nel DEM come il '
    'reticolo di riferimento, e se disponibile un LiDAR a 1 m.')
par('Database (108, 109, 196, 19, 58, 117, 118, 174, 198, 21, 113, 190). Per questi ponti la posizione e\' coerente con il '
    'tracciato del ponte, oppure non esiste nessuna sezione con l\'area del database vicino al ponte: il dato da chiarire e\' '
    'nel database (area riferita a un\'altra opera o a un altro attraversamento, coordinate o progressiva da correggere). '
    'Per 109 e 196 il ponte scavalca chiaramente il corso calcolato; per il tratto SP14 17+200-17+800 (61, 117, 118) le aree '
    'potrebbero essere scambiate tra i ponti.')
titolo('4. Criteri per spostare la sezione di un ponte')
for t in pd.read_excel(xl, sheet_name='Criteri di spostamento').iloc[:, 0]:
    par(str(t))
titolo('5. Prossimi passi')
for t in ['Lanciare il codice con ponti_da_testare.csv sul DEM unico v2 (ponti 6 e 105) in una cartella nuova.',
          'Verificare a vista 107 e 111 con il materiale gia\' preparato.',
          'Reperire il reticolo dei canali per Nova Siri, Scanzano e il 164.',
          'Inviare a chi ha compilato il database le domande per i 12 ponti del gruppo Database.']:
    par('- ' + t)
doc.save(out)
print('ok', out)
