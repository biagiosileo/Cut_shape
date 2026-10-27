"""
Materiale per la verifica a vista delle posizioni proposte (ortofoto / CTR), per i ponti indicati.

Uscite nella cartella indicata:
  shapefile (EPSG:32633) con stile QGIS (.qml con lo stesso nome, applicato in automatico):
    v_punto_attuale, v_punto_proposto, v_spostamento, v_tracciato_ponte, v_bacino_attuale, v_bacino_proposto,
    v_reticolo_locale (reticolo del codice, acc > 4000, entro RAGGIO_M dai ponti), v_db_inizio_fine
  Verifica_ponti.kml        gli stessi layer per Google Earth
  schede/IDxxx.png          scheda per ponte: rilievo ombreggiato del DEM, reticolo, bacini, tracciato, punti
  Verifica_ponti.xlsx       un ponte per riga: dati, cosa guardare, link Google Maps (satellite), Conferma

Uso: python materiale_verifica.py <revisione_non_compatibili.gpkg> <revisione.csv> <indicazioni_revisione.csv>
        <controllo_coordinate.gpkg> <reticolo_idrografico.geojson> <dem5m_unico.tif> <Database.xlsx>
        <out_dir> <id> [<id> ...]
"""
import os
import sys
from xml.sax.saxutils import escape, quoteattr

import geopandas as gpd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import rasterio
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from rasterio.windows import from_bounds
from shapely.geometry import box

gp_rev, csv_rev, csv_ind, gp_cc, ret_path, dem_path, dbx, out = sys.argv[1:9]
ids = [int(i) for i in sys.argv[9:]]
RAGGIO_M = 600.0
os.makedirs(os.path.join(out, 'schede'), exist_ok=True)

lay = {n: gpd.read_file(gp_rev, layer=n) for n in ('punto_attuale', 'punto_proposto', 'spostamento', 'bacino_attuale',
                                                     'bacino_proposto', 'tracciato_ponte')}
lay = {n: g[g.ID.isin(ids)].reset_index(drop=True) for n, g in lay.items()}
R = pd.read_csv(csv_rev).set_index('ID')
I = pd.read_csv(csv_ind).set_index('ID')
k = gpd.read_file(gp_cc, layer='kml_abbinato').set_index('ID').to_crs(32633)
kt = gpd.read_file(gp_cc, layer='kml_tracciati').to_crs(32633).set_index('nome').geometry
ini = gpd.read_file(gp_cc, layer='db_inizio').to_crs(32633)
fin = gpd.read_file(gp_cc, layer='db_fine').to_crs(32633)
db = pd.read_excel(dbx).rename(columns={'Unnamed: 0': 'ID'})
db = db[pd.to_numeric(db.ID, errors='coerce').notna()]
db['ID'] = db.ID.astype(int)
db = db.set_index('ID')
num = lambda c: pd.to_numeric(db[c].astype(str).str.replace(',', '.').str.extract(r'([\d.]+)')[0], errors='coerce')
luce, alveo, lung = num('LUCE TOT (SPALLA/SPALLA - m)'), num('LARGHEZZA ALVEO INCISO (m)'), num(
    'LUNGHEZZA COMPLESSIVA (INCLUSE LE SPALLE - m)')

# tracciato del ponte: KML abbinato per nome se c'e', altrimenti quello del controllo coordinate
tr = []
for i in ids:
    t = kt.get(k.loc[i, 'KML_nome'])
    if t is None:
        g = lay['tracciato_ponte']
        t = g[g.ID == i].geometry.iloc[0] if (g.ID == i).any() else None
    if t is not None:
        tr.append({'ID': i, 'nome_KML': k.loc[i, 'KML_nome'], 'geometry': t})
lay['tracciato_ponte'] = gpd.GeoDataFrame(tr, crs=32633)
dbp = pd.concat([ini[ini.ID.isin(ids)].assign(tipo='inizio DB'), fin[fin.ID.isin(ids)].assign(tipo='fine DB')])
lay['db_inizio_fine'] = gpd.GeoDataFrame(dbp[['ID', 'tipo', 'geometry']], crs=32633).reset_index(drop=True)
zona = gpd.GeoSeries(lay['punto_attuale'].geometry.buffer(RAGGIO_M), crs=32633).union_all()
ret = gpd.read_file(ret_path, bbox=tuple(zona.bounds)).to_crs(32633)
ret = ret[ret.intersects(zona)].copy()
ret['geometry'] = ret.intersection(zona)
lay['reticolo_locale'] = ret.reset_index(drop=True)

STILI = {  # nome: (tipo, colore RGB, spessore/dimensione, extra)
    'punto_attuale': ('marker', '120,120,120', '3.4', 'circle'), 'punto_proposto': ('marker', '255,200,0', '3.8', 'circle'),
    'spostamento': ('line', '150,60,160', '0.6', 'solid'), 'tracciato_ponte': ('line', '255,120,0', '0.9', 'solid'),
    'bacino_attuale': ('fill', '120,120,120', '0.6', 'dash'), 'bacino_proposto': ('fill', '220,30,30', '0.6', 'solid'),
    'reticolo_locale': ('line', '30,110,220', '0.5', 'solid'), 'db_inizio_fine': ('marker', '0,170,170', '2.6', 'triangle')}


def qml(tipo, col, dim, extra, etichetta):
    if tipo == 'marker':
        lyr = (f'<layer class="SimpleMarker" pass="0" enabled="1" locked="0"><Option type="Map">'
               f'<Option name="color" type="QString" value="{col},255"/><Option name="name" type="QString" value="{extra}"/>'
               f'<Option name="size" type="QString" value="{dim}"/><Option name="outline_color" type="QString" value="30,30,30,255"/>'
               f'<Option name="outline_width" type="QString" value="0.3"/></Option></layer>')
    elif tipo == 'line':
        lyr = (f'<layer class="SimpleLine" pass="0" enabled="1" locked="0"><Option type="Map">'
               f'<Option name="line_color" type="QString" value="{col},255"/><Option name="line_width" type="QString" value="{dim}"/>'
               f'<Option name="line_style" type="QString" value="{extra}"/></Option></layer>')
    else:
        lyr = (f'<layer class="SimpleFill" pass="0" enabled="1" locked="0"><Option type="Map">'
               f'<Option name="color" type="QString" value="{col},35"/><Option name="style" type="QString" value="solid"/>'
               f'<Option name="outline_color" type="QString" value="{col},255"/><Option name="outline_width" type="QString" value="{dim}"/>'
               f'<Option name="outline_style" type="QString" value="{extra}"/></Option></layer>')
    lab = ('<labeling type="simple"><settings calloutType="simple"><text-style fieldName="ID" fontSize="9" '
           'textColor="0,0,0,255"><text-buffer bufferDraw="1" bufferSize="0.8" bufferColor="255,255,255,255"/></text-style>'
           '<placement placement="0" dist="1.5"/><rendering scaleVisibility="0"/></settings></labeling>') if etichetta else ''
    return (f"<!DOCTYPE qgis PUBLIC 'http://mrcc.com/qgis.dtd' 'SYSTEM'><qgis version=\"3.28\" "
            f"styleCategories=\"Symbology|Labeling\" labelsEnabled=\"{1 if etichetta else 0}\"><renderer-v2 type=\"singleSymbol\">"
            f"<symbols><symbol name=\"0\" type=\"{tipo}\" alpha=\"1\">{lyr}</symbol></symbols></renderer-v2>{lab}</qgis>")


for n, g in lay.items():
    f = os.path.join(out, f'v_{n}')
    gpd.GeoDataFrame(g, crs=32633).to_file(f + '.shp', encoding='utf-8')
    t, c, d, e = STILI[n]
    open(f + '.qml', 'w').write(qml(t, c, d, e, n in ('punto_attuale', 'punto_proposto')))

# --- KML ---
KC = {'punto_attuale': 'ff787878', 'punto_proposto': 'ff00c8ff', 'spostamento': 'ffa03c96', 'tracciato_ponte': 'ff0078ff',
      'bacino_attuale': 'ff787878', 'bacino_proposto': 'ff1e1edc', 'reticolo_locale': 'ffdc6e1e', 'db_inizio_fine': 'ffaaaa00'}


def coords(g):
    return ' '.join(f'{x:.7f},{y:.7f},0' for x, y in g.coords)


def gk(g):
    if g.geom_type == 'Point':
        return f'<Point><coordinates>{g.x:.7f},{g.y:.7f},0</coordinates></Point>'
    if g.geom_type == 'LineString':
        return f'<LineString><coordinates>{coords(g)}</coordinates></LineString>'
    if g.geom_type == 'Polygon':
        return (f'<Polygon><outerBoundaryIs><LinearRing><coordinates>{coords(g.exterior)}</coordinates></LinearRing>'
                f'</outerBoundaryIs></Polygon>')
    return '<MultiGeometry>' + ''.join(gk(p) for p in g.geoms if not p.is_empty) + '</MultiGeometry>'


kml = ['<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document><name>Verifica posizioni</name>']
for n, c in KC.items():
    kml.append(f'<Style id="{n}"><IconStyle><color>{c}</color><scale>0.9</scale></IconStyle><LineStyle><color>{c}</color>'
               f'<width>2</width></LineStyle><PolyStyle><fill>0</fill></PolyStyle></Style>')
for n in KC:
    g = gpd.GeoDataFrame(lay[n], crs=32633).to_crs(4326)
    if n.startswith('bacino') or n == 'reticolo_locale':
        g['geometry'] = g.geometry.simplify(0.00002)
    kml.append(f'<Folder><name>{n.replace("_", " ")}</name>')
    for _, r in g.iterrows():
        if r.geometry is None or r.geometry.is_empty:
            continue
        nome = str(int(r.ID)) if 'ID' in g and pd.notna(r.get('ID')) else ''
        kml.append(f'<Placemark><name>{nome}</name><styleUrl>#{n}</styleUrl>{gk(r.geometry)}</Placemark>')
    kml.append('</Folder>')
kml.append('</Document></kml>')
open(os.path.join(out, 'Verifica_ponti.kml'), 'w', encoding='utf-8').write(''.join(kml))


# --- cosa guardare ---
def domanda(i):
    r, d = R.loc[i], I.loc[i]
    rel = str(d.Relazione)
    aff = d['Area affluente maggiore (km²)'] if pd.notna(d['Area affluente maggiore (km²)']) else 0
    if rel.startswith('stesso corso'):
        q = (f"Stesso corso, punto proposto {'a valle' if 'valle' in rel else 'a monte'} di {r.Dist_spostamento_m:.0f} m. ")
        if aff > 0 and 'valle' in rel:
            q += (f"Tra i due punti entra un affluente di {aff:g} km², che il punto proposto include: il ponte riceve "
                  "quell'affluente? SI se l'affluente passa sotto il ponte o confluisce a monte; NO se confluisce a valle.")
        elif aff > 0:
            q += (f"Tra i due punti entra un affluente di {aff:g} km², che il punto proposto esclude: SI se l'affluente "
                  "confluisce a valle del ponte (il ponte non lo riceve); NO se passa sotto il ponte o confluisce a monte.")
        else:
            q += "Nessun affluente in mezzo: SI se il ponte e' su questo corso e il punto proposto e' nell'alveo scavalcato."
    else:
        q = ("Corso diverso: quale dei due fossi/corsi (punto grigio o punto giallo) attraversa il ponte sull'ortofoto? "
             "SI se il giallo.")
        if d['Tracciato attraversa corso attuale'] and not d['Tracciato attraversa corso proposto']:
            q += ' Attenzione: il tracciato KML del ponte attraversa il corso del punto attuale (probabile NO).'
    return q


# --- schede PNG ---
with rasterio.open(dem_path) as src:
    for i in ids:
        g0 = lay['punto_attuale'][lay['punto_attuale'].ID == i].geometry.iloc[0]
        g1 = lay['punto_proposto'][lay['punto_proposto'].ID == i].geometry
        pts = [g0] + ([g1.iloc[0]] if len(g1) else [])
        xs, ys = [p.x for p in pts], [p.y for p in pts]
        m = max(250.0, 0.6 * max(max(xs) - min(xs), max(ys) - min(ys)))
        cx, cy = (max(xs) + min(xs)) / 2, (max(ys) + min(ys)) / 2
        b = (cx - m, cy - m, cx + m, cy + m)
        w = from_bounds(*b, transform=src.transform).round_offsets().round_lengths()
        z = src.read(1, window=w).astype(float)
        z[z == src.nodata] = np.nan
        z *= src.scales[0]
        gy, gx = np.gradient(z, 5.0)
        sl, asp = np.pi / 2 - np.arctan(np.hypot(gx, gy)), np.arctan2(-gx, gy)
        az, alt = np.radians(315), np.radians(45)
        hs = np.sin(alt) * np.sin(sl) + np.cos(alt) * np.cos(sl) * np.cos(az - asp)
        Tw = src.window_transform(w)
        ext = [Tw.c, Tw.c + z.shape[1] * 5, Tw.f - z.shape[0] * 5, Tw.f]
        fig, ax = plt.subplots(figsize=(8.3, 8.3))
        ax.imshow(hs, extent=ext, cmap='gray', vmin=0, vmax=1)
        ax.imshow(z, extent=ext, cmap='terrain', alpha=0.25)
        cl = box(*b)
        for n, col, lw, ls in (('bacino_attuale', '#555555', 1.4, '--'), ('bacino_proposto', '#d7191c', 1.4, '-')):
            g = lay[n][lay[n].ID == i]
            if len(g):
                gpd.GeoSeries(g.geometry.intersection(cl)).boundary.plot(ax=ax, color=col, lw=lw, ls=ls)
        rr = lay['reticolo_locale'].clip(cl)
        if len(rr):
            rr.plot(ax=ax, color='#1e6edc', lw=1.2)
        t = lay['tracciato_ponte'][lay['tracciato_ponte'].ID == i]
        if len(t):
            t.plot(ax=ax, color='#ff7800', lw=3)
        dd = lay['db_inizio_fine'][lay['db_inizio_fine'].ID == i].clip(cl)
        if len(dd):
            dd.plot(ax=ax, color='#00aaaa', marker='^', markersize=40, edgecolor='k')
        ax.plot(g0.x, g0.y, 'o', ms=10, mfc='#888888', mec='k')
        if len(g1):
            ax.plot(g1.iloc[0].x, g1.iloc[0].y, 'o', ms=11, mfc='#ffc800', mec='k')
            ax.plot([g0.x, g1.iloc[0].x], [g0.y, g1.iloc[0].y], color='#963ca0', lw=1.2)
        ax.set_xlim(b[0], b[2]); ax.set_ylim(b[1], b[3])
        ax.set_xticks([]); ax.set_yticks([])
        ax.set_xlabel(''); ax.set_ylabel('')
        ax.plot([b[0] + 20, b[0] + 120], [b[1] + 20, b[1] + 20], color='k', lw=3)
        ax.text(b[0] + 70, b[1] + 30, '100 m', ha='center', fontsize=8)
        r = R.loc[i]
        nome = str(db.loc[i, 'Nome opera']).strip()
        nome = '' if nome in ('-', 'nan') else nome + ' '
        ax.set_title(f"ID {i} - {nome}({db.loc[i, 'NUOVA NUMERAZIONE']} km "
                     f"{db.loc[i, 'Progressiva km']})\nArea DB {r.A_db_km2:g} km²  |  attuale (grigio) {r.A_attuale_km2:g}  |  "
                     f"proposta (giallo) {r.A_proposta_km2:g}", fontsize=10)
        fig.text(0.02, 0.01, 'grigio: punto e bacino attuale (tratteggio) - giallo/rosso: punto e bacino proposto - '
                 'arancio: tracciato del ponte (KML) - blu: reticolo - triangoli: inizio/fine DB', fontsize=7)
        fig.savefig(os.path.join(out, 'schede', f'ID{i:03d}.png'), dpi=110, bbox_inches='tight')
        plt.close(fig)

# --- Excel ---
pa = lay['punto_attuale'].set_index('ID').to_crs(4326).geometry
pp = lay['punto_proposto'].set_index('ID').to_crs(4326).geometry
link = lambda p: f'https://www.google.com/maps/@{p.y:.6f},{p.x:.6f},19z/data=!3m1!1e3'
righe = []
for i in ids:
    r, d = R.loc[i], I.loc[i]
    righe.append({'ID': i, 'Nome opera': str(db.loc[i, 'Nome opera']).strip(), 'Strada': db.loc[i, 'NUOVA NUMERAZIONE'],
                  'km': db.loc[i, 'Progressiva km'], 'Area DB (km²)': r.A_db_km2, 'Area attuale (km²)': r.A_attuale_km2,
                  'Area proposta (km²)': r.A_proposta_km2, 'Spostamento (m)': r.Dist_spostamento_m,
                  'Relazione tra i punti': d.Relazione, 'Luce ponte DB (m)': luce.get(i), 'Alveo DB (m)': alveo.get(i),
                  'Indicazione automatica': d.Indicazione, 'Cosa guardare': domanda(i),
                  'Google Maps punto attuale': link(pa[i]), 'Google Maps punto proposto': link(pp[i]),
                  'Scheda': f'schede/ID{i:03d}.png', 'Conferma (SI/NO)': '', 'Note': ''})
T = pd.DataFrame(righe)
xl = os.path.join(out, 'Verifica_ponti.xlsx')
guida = pd.DataFrame({'Come fare la verifica': [
    "1. In QGIS caricare gli shapefile v_*.shp (lo stile si applica da solo) sopra un'ortofoto (es. XYZ Google Satellite "
    "o WMS regionale) e/o la CTR. In alternativa aprire Verifica_ponti.kml in Google Earth.",
    "2. Per ogni ponte guardare la scheda (cartella schede) e la colonna 'Cosa guardare'; i link aprono Google Maps in "
    "vista satellite sul punto attuale e su quello proposto.",
    "3. Il reticolo del DEM (blu) e' una linea di modello: puo' essere spostato di qualche decina di metri rispetto "
    "all'alveo reale. Conta quale corso attraversa il ponte e dove confluiscono gli affluenti rispetto al ponte.",
    "4. SI = il ponte scavalca il corso del punto proposto (giallo) e, se c'e' un affluente in mezzo, l'affluente passa "
    "sotto il ponte o confluisce a monte. NO = il ponte e' sul corso del punto attuale, oppure l'affluente confluisce a "
    "valle del ponte.",
    "5. Compilare 'Conferma (SI/NO)' e rimandare il file: ricalcolo solo i ponti con SI."]})
with pd.ExcelWriter(xl, engine='openpyxl') as w:
    guida.to_excel(w, sheet_name='Come verificare', index=False)
    T.to_excel(w, sheet_name='Ponti da verificare', index=False)
wb = load_workbook(xl)
ws = wb['Ponti da verificare']
L_ = {c: chr(65 + n) for n, c in enumerate(T.columns)}
dv = DataValidation(type='list', formula1='"SI,NO"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"{L_['Conferma (SI/NO)']}2:{L_['Conferma (SI/NO)']}{ws.max_row}")
for r in range(2, ws.max_row + 1):
    for c in ('Google Maps punto attuale', 'Google Maps punto proposto', 'Scheda'):
        cell = ws[f'{L_[c]}{r}']
        cell.hyperlink = cell.value
        cell.value = 'apri' if c != 'Scheda' else cell.value
        cell.font = Font(color='0563C1', underline='single')
    ws[f"{L_['Conferma (SI/NO)']}{r}"].fill = PatternFill('solid', fgColor='FFF2CC')
    for c in ws[r]:
        c.alignment = Alignment(wrap_text=True, vertical='top')
for c, wd in zip(T.columns, [6, 22, 8, 9, 10, 10, 10, 10, 26, 9, 9, 16, 60, 11, 11, 16, 11, 25]):
    ws.column_dimensions[L_[c]].width = wd
for s in wb.worksheets:
    for c in s[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDDDDD')
        c.alignment = Alignment(wrap_text=True, vertical='center')
ws.freeze_panes = 'C2'
wb['Come verificare'].column_dimensions['A'].width = 140
wb.save(xl)
print('ok', len(ids), 'ponti')
