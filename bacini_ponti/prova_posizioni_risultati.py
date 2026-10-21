"""
Risultati di prova con le posizioni proposte (non ancora confermate): confronto con l'area del database per
tutti i 212 ponti, prendendo per ogni ponte il risultato piu' recente:
  prova (posizioni proposte) > correzione delle porzioni > esecuzione dell'utente.

Uso: python prova_posizioni_risultati.py <parametri utente.csv> <parametri porzioni corrette.csv>
        <cartella porzioni della prova> <origine_posizioni.csv> <Database.xlsx> <out_dir>
"""
import glob
import json
import os
import sys

import geopandas as gpd
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from shapely.geometry import shape

pu, pc, cart, orig, dbx, out = sys.argv[1:7]
os.makedirs(out, exist_ok=True)


def leggi(f):
    d = pd.read_csv(f)
    d['ID'] = d.ID.astype(int)
    return d.set_index('ID')


U, C = leggi(pu), leggi(pc)
N = pd.concat([pd.read_csv(f) for f in sorted(glob.glob(os.path.join(cart, '*_parametri.csv')))])
N['ID'] = N.ID.astype(int)
N = N.set_index('ID')
O = pd.read_csv(orig).set_index('ID').origine
db = pd.read_excel(dbx).rename(columns={'Unnamed: 0': 'ID'})
db = db[pd.to_numeric(db.ID, errors='coerce').notna()]
db['ID'] = db.ID.astype(int)
db = db.set_index('ID')
adb = pd.to_numeric(db['SUPERFICIE BACINO (km2)'], errors='coerce')

F = U.copy()
F.update(C)
F.update(N)
fonte = pd.Series('esecuzione utente', index=F.index)
fonte[C.index] = 'porzione corretta'
fonte[N.index] = [f'prova: {O.get(i, "")}' for i in N.index]


def classe(a, b):
    if pd.isna(b):
        return 'senza area DB'
    x = abs(a / b - 1)
    return 'compatibile' if x <= .15 else ('quasi compatibile' if x <= .25 else 'non compatibile')


T = pd.DataFrame({'ID': F.index, 'Nome opera': [str(db.loc[i, 'Nome opera']).strip() for i in F.index],
                  'Area DB (km²)': adb.reindex(F.index).values, 'Area prima (km²)': U.A_km2.reindex(F.index).values,
                  'Area prova (km²)': F.A_km2.values, 'Diff prima': None, 'Diff prova': None,
                  'Classe prima': [classe(U.A_km2[i], adb.get(i)) for i in F.index],
                  'Classe prova': [classe(F.A_km2[i], adb.get(i)) for i in F.index],
                  'Origine risultato': fonte.values}).sort_values('ID')
cambiati = T[T['Origine risultato'] != 'esecuzione utente']
xl = os.path.join(out, 'Risultati_prova_posizioni.xlsx')
with pd.ExcelWriter(xl, engine='openpyxl') as w:
    pd.DataFrame({'Voce': []}).to_excel(w, sheet_name='Riepilogo', index=False)
    T.to_excel(w, sheet_name='Tutti i ponti', index=False)
    cambiati.to_excel(w, sheet_name='Ponti ricalcolati', index=False)
    F.reset_index().to_excel(w, sheet_name='Parametri prova', index=False)
wb = load_workbook(xl)
for nome in ('Tutti i ponti', 'Ponti ricalcolati'):
    ws = wb[nome]
    n = ws.max_row
    for r in range(2, n + 1):
        ws[f'F{r}'] = f'=IF(C{r}="","",D{r}/C{r}-1)'
        ws[f'G{r}'] = f'=IF(C{r}="","",E{r}/C{r}-1)'
        for c in 'FG':
            ws[f'{c}{r}'].number_format = '0.0%'
        for c in 'CDE':
            ws[f'{c}{r}'].number_format = '0.000'
    ws.conditional_formatting.add(f'I2:I{n}', FormulaRule(formula=[f'I2="compatibile"'], fill=PatternFill('solid', fgColor='C6EFCE')))
    ws.conditional_formatting.add(f'I2:I{n}', FormulaRule(formula=[f'I2="quasi compatibile"'], fill=PatternFill('solid', fgColor='DDEBF7')))
    ws.conditional_formatting.add(f'I2:I{n}', FormulaRule(formula=[f'I2="non compatibile"'], fill=PatternFill('solid', fgColor='FCE4D6')))
    for c, wdt in zip('ABCDEFGHIJ', [6, 30, 11, 11, 11, 10, 10, 17, 17, 34]):
        ws.column_dimensions[c].width = wdt
    ws.freeze_panes = 'C2'
ws = wb['Riepilogo']
ws['A1'], ws['B1'], ws['C1'] = 'Classe', 'Prima (esecuzione utente)', 'Prova (posizioni proposte)'
n = wb['Tutti i ponti'].max_row
for r, c in enumerate(['compatibile', 'quasi compatibile', 'non compatibile', 'senza area DB'], start=2):
    ws[f'A{r}'] = c
    ws[f'B{r}'] = f"=COUNTIF('Tutti i ponti'!H2:H{n},A{r})"
    ws[f'C{r}'] = f"=COUNTIF('Tutti i ponti'!I2:I{n},A{r})"
ws['A6'], ws['B6'], ws['C6'] = 'Totale', '=SUM(B2:B5)', '=SUM(C2:C5)'
ws['A8'] = ("Prova: posizioni proposte NON ancora confermate (revisione su ortofoto da fare). "
            "Ponti ricalcolati: foglio 'Ponti ricalcolati'.")
ws.column_dimensions['A'].width = 22
ws.column_dimensions['B'].width = 26
ws.column_dimensions['C'].width = 26
for w_ in wb.worksheets:
    for c in w_[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDDDDD')
        c.alignment = Alignment(wrap_text=True, vertical='center')
wb.save(xl)

feats = []
for f in sorted(glob.glob(os.path.join(cart, '*_bacini_afferenti.geojson'))):
    feats += json.load(open(f))['features']
g = gpd.GeoDataFrame({'ID': [int(ft['properties']['ID']) for ft in feats]},
                     geometry=[shape(ft['geometry']) for ft in feats], crs=32633).sort_values('ID')
g['A_km2'] = (g.area / 1e6).round(3)
g.to_file(os.path.join(out, 'bacini_prova.shp'))
print(T['Classe prima'].value_counts().to_dict())
print(T['Classe prova'].value_counts().to_dict())
print(cambiati.drop(columns='Nome opera').to_string())
