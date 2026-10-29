"""
Ricapitolo del lavoro sui ponti (cartella separata):
  - Ricapitolo_confronto_aree.xlsx: Statistiche e Confronto aree (come Confronto_aree_database_v2, con formule)
    piu' i fogli Ricapitolo correzioni, Non compatibili - azioni, Criteri di spostamento;
  - shapefile dei 20 ponti non compatibili (punto, bacino, tracciato KML, inizio/fine DB, punto proposto,
    taglio del tombino, reticolo locale) con stili QGIS e KML.

Uso: python ricapitolo.py <parametri prima esecuzione.csv> <parametri definitivi.csv> <Database.xlsx>
        <modifiche_rispetto_a_ponti_originale.csv> <ponti.csv finale> <controllo_coordinate.gpkg>
        <revisione_non_compatibili.gpkg> <prova_tombino.gpkg> <reticolo_idrografico.geojson>
        <cartelle porzioni (separate da virgola, in ordine di priorita')> <out_dir> <script excel_confronto_aree_finale.py>
"""
import json
import os
import subprocess
import sys
from xml.sax.saxutils import escape

import geopandas as gpd
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from shapely.geometry import shape

(p_prima, p_fin, dbx, modif, pcsv, gp_cc, gp_rev, gp_tomb, ret_path, cartelle, out, xl_script) = sys.argv[1:13]
os.makedirs(out, exist_ok=True)


def L(f):
    d = pd.read_csv(f)
    d['ID'] = d.ID.astype(int)
    return d.set_index('ID')


O, F = L(p_prima), L(p_fin)
db = pd.read_excel(dbx).rename(columns={'Unnamed: 0': 'ID'})
db = db[pd.to_numeric(db.ID, errors='coerce').notna()]
db['ID'] = db.ID.astype(int)
db = db.set_index('ID')
adb = pd.to_numeric(db['SUPERFICIE BACINO (km2)'], errors='coerce')


def classe(a, b):
    if pd.isna(b):
        return 'senza area DB'
    x, d = abs(a / b - 1), abs(a - b)
    return 'compatibile' if x <= .15 else ('quasi compatibile' if (x <= .25 or (d <= 0.1 and x <= .5)) else 'non compatibile')


nome = lambda i: '' if str(db.loc[i, 'Nome opera']).strip() in ('-', 'nan') else str(db.loc[i, 'Nome opera']).strip()

AZ = {  # ponte: (gruppo, cosa e' emerso, cosa fare)
    6: ('DEM esteso', "Bacino dell'Agri fuori dal DEM unico a ovest di E 576 km: area per difetto (1.137 km²).",
        "DEM unico v2 con l'alto Agri riempito dal DEM Copernicus 30 m: da testare con ponti_da_testare.csv."),
    105: ('Tombino', 'Il rilevato stradale fa da diga nel DEM: il fosso del ponte non passa sotto la strada.',
          'DEM unico v2 con il taglio del tombino e sezione sul taglio: da testare (prova: 0,114 km², -5%).'),
    107: ('Verifica a vista', 'Il ponte e\' sulla confluenza di due rami (6 e 12 m dal ponte); il DB (0,72) indica il ramo piccolo.',
          "Ortofoto: quale dei due rami passa sotto il ponte? Se il ramo piccolo, sezione su quel ramo."),
    111: ('Verifica a vista', "Il tracciato KML e' parallelo al torrente (15,8 km²) a 16 m; con il tombino il fosso locale "
          "ha 0,15 km². Alveo DB 32 m: coerente con il torrente.", "Ortofoto: il ponte scavalca il torrente? Se si', "
          "sezione sul torrente (punto proposto, 15,79 km²)."),
    13: ('Bonifica / canali', 'SP107 Nova Siri, pianura di bonifica: il deflusso dipende dai canali, non dal DEM.',
         'Reticolo dei canali (Consorzio di Bonifica / CTR) da incidere nel DEM; LiDAR 1 m se disponibile.'),
    154: ('Bonifica / canali', 'Come il 13 (stesso tratto della SP107, 30 m di distanza).',
          'Reticolo dei canali da incidere nel DEM; LiDAR 1 m se disponibile.'),
    155: ('Bonifica / canali', 'Ponte sul Dosso, piana di Scanzano: area condizionata dai canali.',
          'Reticolo dei canali da incidere nel DEM; LiDAR 1 m se disponibile.'),
    164: ('Bonifica / canali', 'Reticolo artificiale di scolo (DB): il DEM aggancia il corso naturale (216 km²).',
          'Reticolo dei canali da incidere nel DEM.'),
    108: ('Database', 'Rilevato stradale come diga, ma con il tombino il fosso ha solo 0,017 km² (DB 0,27).',
          "Verificare sull'ortofoto da dove arriva il fosso; chiedere come e' stata calcolata l'area DB."),
    109: ('Database', "Il ponte scavalca il Fosso dell'Acqua Fetente (6,6 km², angolo 80°): posizione corretta.",
          "Chiedere l'area DB (0,31): probabilmente di un'altra opera o di un fosso laterale."),
    196: ('Database', 'Il ponte scavalca il Bilioso (87 km², angolo 79°): posizione corretta. Nella documentazione '
          "l'opera si chiamava Bilioso II km 0+752.", "Chiedere l'area DB (0,83): probabilmente di un'altra opera."),
    19: ('Database', 'Sezione con area DB a 820 m: le coordinate potrebbero indicare un altro attraversamento.',
         'Verificare coordinate/progressiva con chi ha compilato il DB.'),
    58: ('Database', "Sezione con area DB a 838 m; larghezza dell'alveo coerente con l'area calcolata.",
         'Verificare coordinate/progressiva e valore di area con chi ha compilato il DB.'),
    117: ('Database', 'Sezione con area DB a 467 m; tratto SP14 17+200-17+800 (61, 117, 118) da verificare insieme.',
          'Verificare con chi ha compilato il DB le aree dei tre ponti della SP14.'),
    118: ('Database', "Area calcolata (0,751) = area DB del ponte 61 (0,75): possibile scambio nel tratto SP14.",
          'Verificare con chi ha compilato il DB le aree dei tre ponti della SP14.'),
    174: ('Database', 'Sezione con area DB a 1.159 m: le coordinate potrebbero indicare un altro attraversamento.',
          'Verificare coordinate/progressiva con chi ha compilato il DB.'),
    198: ('Database', 'Sezione con area DB a 702 m: le coordinate potrebbero indicare un altro attraversamento.',
          'Verificare coordinate/progressiva con chi ha compilato il DB.'),
    21: ('Database', "Viadotto sulla ferrovia: 1.333 km² incompatibili con un alveo di 17 m; nessun corso con quell'area "
         'entro 2 km.', "Verificare che sia un attraversamento del reticolo e da dove viene l'area DB."),
    113: ('Database', 'Luce 4 m (L < 6 m): 8,49 km² poco plausibili; nessuna sezione compatibile entro 2 km.',
          "Chiedere da dove viene l'area DB."),
    190: ('Database', 'Nessuna sezione compatibile entro 2 km.', "Verificare coordinate e area DB."),
}

# ---------- Excel ----------
xl = os.path.join(out, 'Ricapitolo_confronto_aree.xlsx')
subprocess.run([sys.executable, '-I', xl_script, p_fin, dbx, xl], check=True)
M = pd.read_csv(modif).set_index('id')
ric = []
for i in M.index:
    ric.append({'ID': i, 'Nome opera': nome(i), 'Correzione': M.loc[i, 'motivo'], 'Spostamento (m)': M.loc[i, 'spostamento (m)'],
                'Area DB (km²)': adb.get(i), 'Area prima (km²)': O.A_km2[i], 'Area dopo (km²)': F.A_km2[i],
                'Classe prima': classe(O.A_km2[i], adb.get(i)), 'Classe dopo': classe(F.A_km2[i], adb.get(i))})
RC = pd.DataFrame(ric)
NC = pd.DataFrame([{'ID': i, 'Nome opera': nome(i), 'Strada': db.loc[i, 'NUOVA NUMERAZIONE'], 'km': db.loc[i, 'Progressiva km'],
                    'Area DB (km²)': adb[i], 'Area calcolata (km²)': F.A_km2[i], 'Differenza %': F.A_km2[i] / adb[i] - 1,
                    'Gruppo': g, 'Cosa e\' emerso': e, 'Cosa fare': a} for i, (g, e, a) in AZ.items()])
ordine = ['DEM esteso', 'Tombino', 'Verifica a vista', 'Bonifica / canali', 'Database']
NC = NC.sort_values(['Gruppo', 'ID'], key=lambda s: s.map({g: k for k, g in enumerate(ordine)}) if s.name == 'Gruppo' else s)
CR = pd.DataFrame({'Criteri per spostare la sezione di un ponte': [
    "L'area del database non e' un criterio per spostare la sezione: e' il dato da verificare. Si usa solo come controllo a posteriori.",
    "1. Tracciato del ponte (KML). La sezione deve stare dove il tracciato attraversa il reticolo, entro 10 m (due celle). "
    "Se il punto e' sul tracciato, non si sposta.",
    "2. Geometria del ponte. Il ponte attraversa il corso di traverso: angolo tra tracciato e asta vicino a 90°. Un'asta parallela "
    "al tracciato (angolo < 30°) non e' quella scavalcata dal ponte: di solito e' il fosso al piede del rilevato stradale.",
    "3. Spostamento massimo. Entro meta' della lunghezza del ponte (o della luce, o della larghezza dell'alveo) piu' 20 m, cioe' "
    "dentro l'impronta del ponte. Oltre, serve una prova esterna (ortofoto, CTR, sopralluogo).",
    "4. Confluenze. Se tra il punto attuale e quello proposto entra un affluente, decide la posizione reale della confluenza "
    "rispetto al ponte: l'affluente si conta solo se passa sotto il ponte o confluisce a monte (caso del ponte 72: NO).",
    "5. Coerenza fisica. La larghezza dell'alveo e la luce del ponte (database) devono essere coerenti con l'area del corso "
    "scelto (relazione larghezza-area dei ponti compatibili).",
    "6. Rilevati stradali. Se il DEM non fa passare l'acqua sotto la strada (asta parallela al ponte, area quasi nulla sotto il "
    "ponte), si apre il tombino nel DEM invece di spostare la sezione (caso del ponte 105).",
    "7. Ogni spostamento va confermato sull'ortofoto e annotato (colonna Conferma); i casi che non soddisfano 1-6 restano "
    "con il punto attuale e si chiariscono con il database."]})
with pd.ExcelWriter(xl, engine='openpyxl', mode='a') as w:
    RC.to_excel(w, sheet_name='Ricapitolo correzioni', index=False)
    NC.to_excel(w, sheet_name='Non compatibili - azioni', index=False)
    CR.to_excel(w, sheet_name='Criteri di spostamento', index=False)
wb = load_workbook(xl)
st = wb['Statistiche']
r0 = 26
st[f'A{r0}'] = 'Ricapitolo (dalla prima esecuzione al run definitivo)'
st[f'A{r0}'].font = Font(bold=True, size=11)
tab = [('Classe', 'Prima esecuzione', 'Run definitivo')] + [
    (c, int(sum(classe(O.A_km2[i], adb.get(i)) == c for i in O.index)), int(sum(classe(F.A_km2[i], adb.get(i)) == c for i in F.index)))
    for c in ('compatibile', 'quasi compatibile', 'non compatibile', 'senza area DB')]
for k, row in enumerate(tab):
    for j, v in enumerate(row):
        st.cell(r0 + 1 + k, 1 + j, v).font = Font(bold=(k == 0))
st[f'A{r0 + 7}'] = 'Ponti modificati rispetto al ponti.csv originale'
st[f'C{r0 + 7}'] = f"=COUNTA('Ricapitolo correzioni'!A2:A{len(RC) + 1})"
st[f'A{r0 + 8}'] = 'Non compatibili per gruppo di azione'
st[f'A{r0 + 8}'].font = Font(bold=True)
for k, g in enumerate(ordine):
    st[f'A{r0 + 9 + k}'] = g
    st[f'C{r0 + 9 + k}'] = f"=COUNTIF('Non compatibili - azioni'!H2:H{len(NC) + 1},A{r0 + 9 + k})"
ws = wb['Ricapitolo correzioni']
for r in range(2, ws.max_row + 1):
    for c in ws[r]:
        c.alignment = Alignment(wrap_text=True, vertical='top')
for c, wd in zip('ABCDEFGHI', [6, 24, 60, 11, 11, 11, 11, 16, 16]):
    ws.column_dimensions[c].width = wd
ws = wb['Non compatibili - azioni']
for r in range(2, ws.max_row + 1):
    ws[f'G{r}'].number_format = '+0%;-0%;0%'
    for c in ws[r]:
        c.alignment = Alignment(wrap_text=True, vertical='top')
for c, wd in zip('ABCDEFGHIJ', [6, 24, 8, 9, 10, 10, 10, 18, 60, 60]):
    ws.column_dimensions[c].width = wd
wb['Criteri di spostamento'].column_dimensions['A'].width = 150
for r in range(2, wb['Criteri di spostamento'].max_row + 1):
    wb['Criteri di spostamento'][f'A{r}'].alignment = Alignment(wrap_text=True, vertical='top')
for s in ('Ricapitolo correzioni', 'Non compatibili - azioni', 'Criteri di spostamento'):
    for c in wb[s][1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDDDDD')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    wb[s].freeze_panes = 'A2' if s == 'Criteri di spostamento' else 'C2'
wb.save(xl)
print('Excel ok')

# ---------- shapefile dei non compatibili ----------
ids = list(AZ)
P = pd.read_csv(pcsv).set_index('id')
pa = gpd.GeoDataFrame({'ID': ids}, geometry=gpd.points_from_xy(P.X[ids], P.Y[ids]), crs=4326).to_crs(32633)
pa['nome'] = [nome(i) for i in ids]
pa['gruppo'] = [AZ[i][0] for i in ids]
pa['A_db'] = [adb[i] for i in ids]
pa['A_calc'] = [F.A_km2[i] for i in ids]
pa['diff_pct'] = [round(100 * (F.A_km2[i] / adb[i] - 1), 1) for i in ids]
pa['azione'] = [AZ[i][2][:250] for i in ids]
feats = {}
for cart in cartelle.split(','):
    for f in sorted(os.listdir(cart)):
        if f.endswith('_bacini_afferenti.geojson'):
            for ft in json.load(open(os.path.join(cart, f)))['features']:
                i = int(ft['properties']['ID'])
                if i in ids and i not in feats:
                    feats[i] = shape(ft['geometry'])
bac = gpd.GeoDataFrame({'ID': list(feats), 'gruppo': [AZ[i][0] for i in feats], 'A_km2': [F.A_km2[i] for i in feats]},
                       geometry=list(feats.values()), crs=32633)
k = gpd.read_file(gp_cc, layer='kml_abbinato').set_index('ID').to_crs(32633)
kt = gpd.read_file(gp_cc, layer='kml_tracciati').to_crs(32633).set_index('nome').geometry
tu = gpd.read_file(gp_cc, layer='tracciato_usato').set_index('ID').to_crs(32633)
tr = gpd.GeoDataFrame([{'ID': i, 'fonte': 'KML per nome' if kt.get(k.loc[i, 'KML_nome']) is not None else tu.fonte[i],
                        'geometry': kt.get(k.loc[i, 'KML_nome']) if kt.get(k.loc[i, 'KML_nome']) is not None else tu.geometry[i]}
                       for i in ids], crs=32633)
ini = gpd.read_file(gp_cc, layer='db_inizio').to_crs(32633)
fin = gpd.read_file(gp_cc, layer='db_fine').to_crs(32633)
dbp = gpd.GeoDataFrame(pd.concat([ini[ini.ID.isin(ids)].assign(tipo='inizio DB'), fin[fin.ID.isin(ids)].assign(tipo='fine DB')])
                       [['ID', 'tipo', 'geometry']], crs=32633)
pp = gpd.read_file(gp_rev, layer='punto_proposto')
pp = pp[pp.ID.isin([107, 111])][['ID', 'A_km2', 'geometry']].assign(origine='revisione (sezione compatibile vicina)')
tl = gpd.read_file(gp_tomb, layer='linee')
tl = tl[(tl.tipo == 'taglio tombino') & tl.ID.isin([105])]
zona = pa.geometry.buffer(600).union_all()
ret = gpd.read_file(ret_path, bbox=tuple(zona.bounds)).to_crs(32633)
ret = ret[ret.intersects(zona)].copy()
ret['geometry'] = ret.intersection(zona)
LAY = {'nc_punto': (pa, 'marker', '215,25,28', '3.6', 'circle', True), 'nc_bacino': (bac, 'fill', '215,25,28', '0.6', 'solid', False),
       'nc_tracciato_ponte': (tr, 'line', '255,120,0', '0.9', 'solid', False),
       'nc_db_inizio_fine': (dbp, 'marker', '0,170,170', '2.6', 'triangle', False),
       'nc_punto_proposto': (pp, 'marker', '255,200,0', '3.6', 'circle', True),
       'nc_taglio_tombino': (tl, 'line', '30,110,220', '0.8', 'solid', False),
       'nc_reticolo_locale': (ret.reset_index(drop=True), 'line', '30,110,220', '0.4', 'solid', False)}


def qml(tipo, col, dim, extra, etichetta):
    if tipo == 'marker':
        lyr = (f'<layer class="SimpleMarker" pass="0" enabled="1" locked="0"><Option type="Map"><Option name="color" '
               f'type="QString" value="{col},255"/><Option name="name" type="QString" value="{extra}"/><Option name="size" '
               f'type="QString" value="{dim}"/><Option name="outline_color" type="QString" value="30,30,30,255"/></Option></layer>')
    elif tipo == 'line':
        lyr = (f'<layer class="SimpleLine" pass="0" enabled="1" locked="0"><Option type="Map"><Option name="line_color" '
               f'type="QString" value="{col},255"/><Option name="line_width" type="QString" value="{dim}"/></Option></layer>')
    else:
        lyr = (f'<layer class="SimpleFill" pass="0" enabled="1" locked="0"><Option type="Map"><Option name="color" '
               f'type="QString" value="{col},30"/><Option name="outline_color" type="QString" value="{col},255"/>'
               f'<Option name="outline_width" type="QString" value="{dim}"/></Option></layer>')
    lab = ('<labeling type="simple"><settings calloutType="simple"><text-style fieldName="ID" fontSize="9" '
           'textColor="0,0,0,255"><text-buffer bufferDraw="1" bufferSize="0.8" bufferColor="255,255,255,255"/></text-style>'
           '<placement placement="0" dist="1.5"/></settings></labeling>') if etichetta else ''
    return (f"<!DOCTYPE qgis PUBLIC 'http://mrcc.com/qgis.dtd' 'SYSTEM'><qgis version=\"3.28\" styleCategories=\"Symbology|"
            f"Labeling\" labelsEnabled=\"{1 if etichetta else 0}\"><renderer-v2 type=\"singleSymbol\"><symbols><symbol "
            f"name=\"0\" type=\"{tipo}\" alpha=\"1\">{lyr}</symbol></symbols></renderer-v2>{lab}</qgis>")


kml = ['<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document><name>Non compatibili</name>']
for n, (g, t, col, dim, ex, et) in LAY.items():
    g = gpd.GeoDataFrame(g, crs=32633)
    g.to_file(os.path.join(out, n + '.shp'), encoding='utf-8')
    open(os.path.join(out, n + '.qml'), 'w').write(qml(t, col, dim, ex, et))
    gg = g.to_crs(4326)
    if t != 'marker':
        gg['geometry'] = gg.geometry.simplify(0.00002)
    kml.append(f'<Folder><name>{n}</name>')
    for _, r in gg.iterrows():
        if r.geometry is None or r.geometry.is_empty:
            continue
        nm = str(int(r.ID)) if 'ID' in gg and pd.notna(r.get('ID')) else ''
        geoms = getattr(r.geometry, 'geoms', [r.geometry])
        parti = []
        for q in geoms:
            if q.geom_type == 'Point':
                parti.append(f'<Point><coordinates>{q.x:.7f},{q.y:.7f},0</coordinates></Point>')
            elif q.geom_type == 'LineString':
                parti.append('<LineString><coordinates>' + ' '.join(f'{x:.7f},{y:.7f},0' for x, y in q.coords) + '</coordinates></LineString>')
            elif q.geom_type == 'Polygon':
                parti.append('<Polygon><outerBoundaryIs><LinearRing><coordinates>' + ' '.join(
                    f'{x:.7f},{y:.7f},0' for x, y in q.exterior.coords) + '</coordinates></LinearRing></outerBoundaryIs></Polygon>')
        desc = escape(f"{r.get('gruppo', '')} {r.get('azione', '')}".strip())
        kml.append(f'<Placemark><name>{nm}</name><description>{desc}</description><MultiGeometry>{"".join(parti)}'
                   f'</MultiGeometry></Placemark>')
    kml.append('</Folder>')
kml.append('</Document></kml>')
open(os.path.join(out, 'Non_compatibili.kml'), 'w', encoding='utf-8').write(''.join(kml))
print('shapefile ok:', {n: len(v[0]) for n, v in LAY.items()})
