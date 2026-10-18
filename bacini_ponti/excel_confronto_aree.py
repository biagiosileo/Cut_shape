"""
Excel di confronto aree DB / codice (pysheds_bridges_5m.py) con tre fogli:
  1 Riepilogo           conteggi (formule COUNTIF sul foglio 2)
  2 Confronto aree      ID, aree, differenza %, classe, valutazione, motivazione
                        (celeste = quasi compatibile 15-25%, arancione = non compatibile >25%)
  3 Parametri morfologici  ID, nome ponte e tutte le variabili prodotte dal codice

Uso: python excel_confronto_aree.py <parametri_codice.csv> <Database.xlsx>
        <Giustificazione_non_compatibili.xlsx> <out.xlsx>
"""
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

pcsv, dbx, gx, out = sys.argv[1:5]
U = pd.read_csv(pcsv)
U['ID'] = U.ID.astype(int)
U = U.set_index('ID').sort_index()
db = pd.read_excel(dbx).rename(columns={'Unnamed: 0': 'ID'})
db = db[pd.to_numeric(db.ID, errors='coerce').notna()]
db['ID'] = db.ID.astype(int)
db = db.set_index('ID')
G = pd.read_excel(gx, sheet_name='Tutti i non compatibili (39)').set_index('ID')
E = pd.read_excel(gx, sheet_name='Errori nel database (19)').set_index('ID')
A_DB = pd.to_numeric(db['SUPERFICIE BACINO (km2)'], errors='coerce')
NOME = db['Nome opera'].astype(str).str.strip()
STRADA = db['NUOVA NUMERAZIONE'].astype(str).str.strip()
PROGR = db['Progressiva km'].astype(str).str.strip()


def nome(i):
    n = NOME.get(i, '')
    return n if n not in ('-', 'nan', '') else f'{STRADA.get(i, "")} km {PROGR.get(i, "")} (senza nome nel DB)'


def motivo(i):
    if i not in G.index:
        return '', ''
    g = G.loc[i]
    cat = g.Categoria
    d = g['Distanza della nuova sezione dal ponte (m)']
    corso = g['Corso della nuova sezione'] if isinstance(g['Corso della nuova sezione'], str) else ''
    if cat == 'Probabile errore database':
        return 'Probabile errore nel database', E.loc[i, 'Giustificazione']
    if cat == 'Area DB sotto soglia':
        return 'Area DB sotto la soglia del reticolo', (
            f"Area DB {g['Area database (km²)']:g} km², inferiore alla soglia del reticolo (0,1 km²): il fosso "
            'attraversato e\' troppo piccolo per il reticolo a 5 m e il codice chiude sul corso minimo '
            'rappresentato. Spostare il punto non serve: l\'area DB e\' plausibile, la differenza dipende dalla soglia.')
    if corso.startswith('stesso'):
        return ('Possibile posizionamento della sezione' if cat == 'Spostamento <=50 m' else
                'Sezione compatibile vicina, da verificare'), (
            f'Sezione con area compatibile a {d:.0f} m dal ponte sullo stesso corso ({corso.split(", ")[1]}): '
            'probabile imprecisione nella posizione della sezione, non errore nel database. Lo spostamento non e\' '
            'adottato automaticamente perche\' la sezione di calcolo resta l\'intersezione tracciato-reticolo; '
            'si puo\' adottare dopo verifica su ortofoto.')
    altro = 'un affluente' if corso.startswith('affluente') else 'un altro corso'
    extra = (' Corso: alveo abbandonato del Bradano Vecchio (aste ISPRA).' if i in (107, 108) else '')
    return 'Ponte forse su un altro corso, da verificare', (
        f'Sezione con area compatibile a {d:.0f} m dal ponte, ma su {altro}: il ponte potrebbe attraversare quel '
        'corso e non quello agganciato. Va verificato su cartografia quale corso attraversa il ponte; spostare '
        'il punto senza questa verifica significherebbe calcolare il bacino di un altro corso.' + extra)


ARIAL = 'Arial'
F = Font(name=ARIAL, size=10)
FB = Font(name=ARIAL, size=10, bold=True)
FT = Font(name=ARIAL, size=13, bold=True)
HEAD = PatternFill('solid', fgColor='D9D9D9')
CELESTE = PatternFill('solid', fgColor='BDD7EE')
ARANCIO = PatternFill('solid', fgColor='F8CBAD')
thin = Side(style='thin', color='BFBFBF')
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()

# ---------------- foglio 2: confronto aree ----------------
ws2 = wb.active
ws2.title = 'Confronto aree'
S2 = "'Confronto aree'"
h2 = ['ID', 'Area database (km²)', 'Area da codice (km²)', 'Differenza %', 'Classe', 'Valutazione', 'Motivazione']
ws2.append(h2)
ids = list(U.index)
for r, i in enumerate(ids, start=2):
    adb = A_DB.get(i)
    val, mot = motivo(i)
    ws2.cell(r, 1, i)
    ws2.cell(r, 2, None if pd.isna(adb) else float(adb))
    ws2.cell(r, 3, float(U.loc[i, 'A_km2']))
    ws2.cell(r, 4, f'=IF(B{r}="","",(C{r}-B{r})/B{r})')
    ws2.cell(r, 5, f'=IF(B{r}="","Senza area DB",IF(ABS(D{r})<=0.15,"Compatibile",'
                   f'IF(ABS(D{r})<=0.25,"Quasi compatibile","Non compatibile")))')
    ws2.cell(r, 6, val)
    ws2.cell(r, 7, mot)
last = len(ids) + 1
for row in ws2.iter_rows(min_row=1, max_row=last, max_col=len(h2)):
    for c in row:
        c.font = FB if c.row == 1 else F
        c.border = BORD
        if c.row == 1:
            c.fill = HEAD
            c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        elif c.column == 7:
            c.alignment = Alignment(wrap_text=True, vertical='top')
        else:
            c.alignment = Alignment(vertical='top')
        if c.row > 1 and c.column in (2, 3):
            c.number_format = '#,##0.000'
        if c.row > 1 and c.column == 4:
            c.number_format = '0.0%'
for col, w in zip('ABCDEFG', (7, 14, 14, 12, 18, 34, 110)):
    ws2.column_dimensions[col].width = w
ws2.row_dimensions[1].height = 30
ws2.freeze_panes = 'B2'
ws2.auto_filter.ref = f'A1:G{last}'
rng = f'A2:G{last}'
ws2.conditional_formatting.add(rng, FormulaRule(formula=[f'$E2="Quasi compatibile"'], fill=CELESTE))
ws2.conditional_formatting.add(rng, FormulaRule(formula=[f'$E2="Non compatibile"'], fill=ARANCIO))

# ---------------- foglio 1: riepilogo ----------------
ws1 = wb.create_sheet('Riepilogo', 0)
ws1['A1'] = 'Confronto tra area del bacino da database e area calcolata dal codice (pysheds_bridges_5m.py)'
ws1['A1'].font = FT
ws1['A2'] = 'Tolleranza di compatibilità: ±15% sull\'area del database. Tutti i conteggi sono formule sul foglio "Confronto aree".'
ws1['A2'].font = F
righe = [
    ('Ponti calcolati dal codice', f'=COUNT({S2}!A2:A{last})', None),
    ('Ponti con area nel database', f'=COUNT({S2}!B2:B{last})', None),
    ('Ponti senza area nel database', f'=COUNTIF({S2}!E2:E{last},"Senza area DB")', None),
    (None, None, None),
    ('Compatibili (differenza ≤ 15%)', f'=COUNTIF({S2}!E2:E{last},"Compatibile")', 'pct'),
    ('Quasi compatibili (15% - 25%)', f'=COUNTIF({S2}!E2:E{last},"Quasi compatibile")', 'pct'),
    ('Non compatibili (> 25%)', f'=COUNTIF({S2}!E2:E{last},"Non compatibile")', 'pct'),
    ('Totale con area nel database', '=SUM(B9:B11)', 'pct'),
]
ws1['A4'], ws1['B4'], ws1['C4'] = 'Voce', 'N. ponti', '% sui ponti con area DB'
r = 5
for lab, f, p in righe:
    if lab:
        ws1.cell(r, 1, lab)
        ws1.cell(r, 2, f)
        if p:
            ws1.cell(r, 3, f'=B{r}/$B$6')
            ws1.cell(r, 3).number_format = '0.0%'
    r += 1
# righe: 5 calcolati, 6 con area, 7 senza, 9 comp, 10 quasi, 11 non, 12 totale
ws1['A14'] = 'Ponti quasi compatibili e non compatibili per valutazione'
ws1['A14'].font = FB
ws1['A15'], ws1['B15'], ws1['C15'], ws1['D15'] = 'Valutazione', 'Quasi compatibili', 'Non compatibili', 'Totale'
vals = ['Probabile errore nel database', 'Area DB sotto la soglia del reticolo',
        'Ponte forse su un altro corso, da verificare', 'Sezione compatibile vicina, da verificare',
        'Possibile posizionamento della sezione']
for k, v in enumerate(vals, start=16):
    ws1.cell(k, 1, v)
    ws1.cell(k, 2, f'=COUNTIFS({S2}!F2:F{last},A{k},{S2}!E2:E{last},"Quasi compatibile")')
    ws1.cell(k, 3, f'=COUNTIFS({S2}!F2:F{last},A{k},{S2}!E2:E{last},"Non compatibile")')
    ws1.cell(k, 4, f'=B{k}+C{k}')
kt = 16 + len(vals)
ws1.cell(kt, 1, 'Totale')
for c in 'BCD':
    ws1[f'{c}{kt}'] = f'=SUM({c}16:{c}{kt - 1})'
note = kt + 2
legenda = [
    'Legenda del foglio "Confronto aree": celeste = quasi compatibile (15-25%), arancione = non compatibile (>25%).',
    'Probabile errore nel database: entro 200 m dal ponte nessuna sezione del reticolo ha area compatibile; spostare il '
    'punto significherebbe calcolare il bacino di un altro punto o di un altro corso. Il tipo di errore (area o coordinate) '
    'è indicato nella motivazione, dal confronto con la larghezza dell\'alveo riportata nel database.',
    'Area DB sotto la soglia del reticolo: fosso più piccolo di 0,1 km², non rappresentato dal reticolo a 5 m.',
    'Ponte forse su un altro corso / sezione compatibile vicina / possibile posizionamento: sezione con area compatibile '
    'entro 200 m, da verificare su ortofoto prima di spostare il punto.',
    'Fonte aree database: colonna "SUPERFICIE BACINO (km2)" di Database_nuovo_11_08_2026.xlsx. '
    'Fonte aree da codice: NEW_parametri_morfologici.csv (lancio dell\'utente).',
]
for k, t in enumerate(legenda):
    ws1.cell(note + k, 1, t)
    ws1.cell(note + k, 1).alignment = Alignment(wrap_text=True, vertical='top')
    ws1.merge_cells(start_row=note + k, start_column=1, end_row=note + k, end_column=4)
    ws1.row_dimensions[note + k].height = 30 if len(t) < 130 else 45
for row in ws1.iter_rows(min_row=3, max_row=note + len(legenda)):
    for c in row:
        if c.font != FT:
            c.font = FB if c.row in (4, 15, 12, kt) or (c.row == 14) else F
for row in list(ws1.iter_rows(min_row=4, max_row=12, max_col=3)) + list(ws1.iter_rows(min_row=15, max_row=kt, max_col=4)):
    for c in row:
        if c.value is not None or c.row in (4, 15):
            c.border = BORD
        if c.row in (4, 15):
            c.fill = HEAD
ws1['A10'].fill = CELESTE
ws1['A11'].fill = ARANCIO
ws1.column_dimensions['A'].width = 52
for c in 'BCD':
    ws1.column_dimensions[c].width = 22

# ---------------- foglio 3: parametri morfologici ----------------
ws3 = wb.create_sheet('Parametri morfologici')
cols = [c for c in U.columns]
ws3.append(['ID', 'Nome ponte'] + cols)
for i in ids:
    vals_ = [U.loc[i, c] for c in cols]
    vals_ = [bool(v) if isinstance(v, (bool,)) or str(v) in ('True', 'False') else
             (None if pd.isna(v) else float(v)) for v in vals_]
    ws3.append([i, nome(i)] + vals_)
for row in ws3.iter_rows(min_row=1, max_row=len(ids) + 1, max_col=len(cols) + 2):
    for c in row:
        c.font = FB if c.row == 1 else F
        c.border = BORD
        if c.row == 1:
            c.fill = HEAD
            c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
ws3.column_dimensions['A'].width = 7
ws3.column_dimensions['B'].width = 46
for k in range(3, len(cols) + 3):
    ws3.cell(1, k).column_letter
    ws3.column_dimensions[ws3.cell(1, k).column_letter].width = 15
ws3.row_dimensions[1].height = 45
ws3.freeze_panes = 'C2'
ws3.auto_filter.ref = ws3.dimensions
wb.save(out)
print('salvato', out)
