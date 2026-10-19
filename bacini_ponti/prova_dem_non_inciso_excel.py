"""
Excel della prova DEM inciso / non inciso sui ponti fuori tolleranza.

Uso: python prova_dem_non_inciso_excel.py <prova_dem_non_inciso.csv> <out.xlsx>
"""
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

d = pd.read_csv(sys.argv[1])
w = d.pivot(index='ID', columns='DEM', values=['A_km2', 'snap_m', 'dist_compat_m', 'A_compat_km2'])
w.columns = [f'{a}_{b}' for a, b in w.columns]
w['A_db'] = d.groupby('ID').A_db_km2.first()


def classe(a, b):
    if pd.isna(b):
        return 'senza area DB'
    x = abs(a / b - 1)
    return 'compatibile' if x <= .15 else ('quasi compatibile' if x <= .25 else 'non compatibile')


def esito(r):
    v = r.A_km2_non_inciso / r.A_km2_inciso - 1
    if abs(v) <= 0.05:
        return 'Nessun effetto: stesso bacino con e senza incisione.'
    if r.snap_m_non_inciso > r.snap_m_inciso + 2:
        return (f'Senza incisione il punto si aggancia a un ramo vicino ({r.snap_m_non_inciso:.0f} m) con area '
                f'diversa ({v:+.0%}); la classe non migliora.')
    return f'Area diversa senza incisione ({v:+.0%}); la classe non migliora.'


T = pd.DataFrame({
    'ID': w.index, 'Area DB (km²)': w.A_db.values,
    'Area DEM inciso (km²)': w.A_km2_inciso.values, 'Area DEM non inciso (km²)': w.A_km2_non_inciso.values,
    'Variazione non inciso / inciso': None, 'Classe DEM inciso': [classe(a, b) for a, b in zip(w.A_km2_inciso, w.A_db)],
    'Classe DEM non inciso': [classe(a, b) for a, b in zip(w.A_km2_non_inciso, w.A_db)],
    'Snap non inciso (m)': w.snap_m_non_inciso.values,
    'Sezione compatibile entro 300 m - inciso (m)': w.dist_compat_m_inciso.values,
    'Sezione compatibile entro 300 m - non inciso (m)': w.dist_compat_m_non_inciso.values,
    'Esito': [esito(r) for _, r in w.iterrows()]})
T.to_excel(sys.argv[2], sheet_name='Inciso vs non inciso', index=False)
wb = load_workbook(sys.argv[2])
ws = wb.active
for r in range(2, ws.max_row + 1):
    ws[f'E{r}'] = f'=D{r}/C{r}-1'
    ws[f'E{r}'].number_format = '0.0%'
    for c in 'BCD':
        ws[f'{c}{r}'].number_format = '0.000'
    ws[f'K{r}'].alignment = Alignment(wrap_text=True, vertical='top')
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill('solid', fgColor='DDDDDD')
    c.alignment = Alignment(wrap_text=True, vertical='center')
for c, wdt in zip('ABCDEFGHIJK', [6, 11, 12, 12, 12, 16, 16, 10, 14, 14, 70]):
    ws.column_dimensions[c].width = wdt
ws.freeze_panes = 'B2'
r = ws.max_row + 2
ws[f'A{r}'] = 'Ponti con classe migliorata senza incisione'
ws[f'G{r}'] = f'=SUMPRODUCT((G2:G{r - 2}="compatibile")*(F2:F{r - 2}<>"compatibile"))'
ws[f'A{r + 1}'] = 'Ponti con area invariata (entro 5%)'
ws[f'G{r + 1}'] = f'=SUMPRODUCT((ABS(E2:E{r - 2})<=0.05)*1)'
wb.save(sys.argv[2])
print(T.drop(columns='Esito').to_string())
