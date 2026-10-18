"""
Excel di verifica del codice sul campione (formule dove possibile):
  Legenda              come si usa e formule
  Grandezze primarie   codice vs ricalcolo indipendente; area e rilievo del ricalcolo sono formule sui dati
                       di base (n. celle, z_max, z_min); differenze e conteggio dei ponti entro tolleranza
                       sono formule
  Parametri derivati   grandezze in ingresso collegate al foglio precedente (sovrascrivibili a mano) e
                       indici calcolati con le stesse formule del codice, confrontati con il valore del codice
  Metodo indipendente  bacino A (codice) vs bacino B: area comune, IoU, aree esclusive, scarti (formule)

Uso: python relazione_excel_verifica.py <cartella dati> <parametri_codice.csv> <Database.xlsx> <out.xlsx>
"""
import glob
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as COL

dati, pcsv, dbx, out = sys.argv[1:5]
V = pd.concat([pd.read_csv(f) for f in glob.glob(os.path.join(dati, 'valori_*.csv'))]).drop_duplicates('ID').set_index('ID')
U = pd.read_csv(pcsv)
U['ID'] = U.ID.astype(int)
U = U.set_index('ID')
db = pd.read_excel(dbx).rename(columns={'Unnamed: 0': 'ID'})
db = db[pd.to_numeric(db.ID, errors='coerce').notna()]
db['ID'] = db.ID.astype(int)
db = db.set_index('ID')
ORD = [2, 1, 24, 167, 72, 106, 90, 174, 129, 154, 107]
ids = [i for i in ORD if i in V.index]
N = len(ids)
CASO = {2: 'compatibile, bacino piccolo', 1: 'compatibile, bacino medio', 24: 'compatibile, bacino grande',
        167: 'compatibile, bacino molto grande (Bradano)', 72: 'quasi compatibile, sezione spostabile < 50 m',
        106: 'quasi compatibile, bacino molto piccolo', 90: 'non compatibile, errore coordinate DB',
        174: 'non compatibile, errore valore area DB', 129: 'area DB sotto soglia del reticolo',
        154: 'area di bonifica piatta', 107: 'altro corso (Bradano Vecchio)'}

AR = 'Arial'
F, FB = Font(name=AR, size=10), Font(name=AR, size=10, bold=True)
FBLU = Font(name=AR, size=10, color='0000FF')      # valori inseriti
FVERDE = Font(name=AR, size=10, color='008000')    # collegamenti ad altri fogli
HEAD = PatternFill('solid', fgColor='D9D9D9')
GIALLO = PatternFill('solid', fgColor='FFF2CC')
VERDE = PatternFill('solid', fgColor='E2EFDA')
ROSSO = PatternFill('solid', fgColor='F8CBAD')
thin = Side(style='thin', color='BFBFBF')
B = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()


def intest(ws, row, labels, widths=None):
    for k, t in enumerate(labels, 1):
        c = ws.cell(row, k, t)
        c.font, c.fill, c.border = FB, HEAD, B
        c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 48
    if widths:
        for k, w in enumerate(widths, 1):
            ws.column_dimensions[COL(k)].width = w


def cella(ws, r, c, v, fmt=None, font=F, fill=None):
    x = ws.cell(r, c, v)
    x.font, x.border = font, B
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    return x


# ---------------- Legenda ----------------
ws = wb.active
ws.title = 'Legenda'
testo = [
    ('Verifica del codice pysheds_bridges_5m.py su un campione di ponti', FB),
    ('', F),
    ('Colori del testo: blu = valore inserito (dal codice o dal ricalcolo); nero = formula; verde = collegamento a un '
     'altro foglio. Celle gialle: grandezze in ingresso che si possono sovrascrivere a mano.', F),
    ('Grandezze primarie: valori del codice (file NEW_parametri_morfologici.csv) e ricalcolo indipendente sullo stesso '
     'bacino. Area = n. celle × 25 m² e rilievo = z_max − z_min sono formule sui dati di base (colonne a destra); gli altri '
     'valori del ricalcolo sono statistiche sul DEM e sul CN calcolate fuori dal foglio. Differenze e conteggi sono formule.', F),
    ('Parametri derivati: le grandezze in ingresso (celle gialle) sono collegate al foglio "Grandezze primarie"; si possono '
     'sovrascrivere con un valore, una alla volta, e gli indici si ricalcolano subito con le formule del codice.', F),
    ('Metodo indipendente: bacino del codice (A) e bacino ricavato con il metodo indipendente descritto nella relazione (B). '
     'IoU = area comune / (A + B − area comune).', F),
    ('', F),
    ('Formule del codice (A in km², L in km, R in m, P in km, AE e z_min in m):', FB),
    ('FF = A / L²   ·   Melton = (R/1000) / √A   ·   HI = (AE − z_min) / R   ·   Re = 2·√(A/π) / L', F),
    ('Shape factor = L / √A   ·   Kc = P / (2·√(π·A))   ·   Rc = 4·π·A / P²   ·   Rr = (R/1000) / L   ·   '
     'Rilievo relativo = R / (P·1000)', F),
]
for k, (t, f) in enumerate(testo, 1):
    c = ws.cell(k, 1, t)
    c.font = f
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[k].height = 44 if len(t) > 200 else (30 if len(t) > 110 else 16)
ws.column_dimensions['A'].width = 140

# ---------------- Grandezze primarie ----------------
ws = wb.create_sheet('Grandezze primarie')
SP = "'Grandezze primarie'"
# (colonna codice, colonna ricalcolo, nome, formato, tolleranza = mezza unita' dell'arrotondamento del codice)
G = [('A_km2', 'A_km2', 'Area A (km²)', '0.0000', 0.0005), ('S', 'S', 'Pendenza media S (-)', '0.00000', 5e-6),
     ('AE_m', 'AE_m', 'Quota media AE (m)', '0.00', 0.005), ('Relief_m', 'Relief_m', 'Rilievo R (m)', '0.00', 0.005),
     ('Longest_flow_length_km', 'L_km', 'Percorso più lungo L (km)', '0.000', 0.0005),
     ('Basin_Perimeter_km', 'P_km', 'Perimetro P (km)', '0.000', 0.0005),
     ('Curve_Number_mean', 'CN_medio', 'CN medio (-)', '0.00', 0.005),
     ('Cell_Slope', 'Cell_Slope', 'Pendenza della cella di sezione (-)', '0.00000', 5e-6),
     ('Reach_Slope_Centered', 'Reach_Slope', 'Pendenza del tratto 500 m (-)', '0.00000', 5e-6)]
lab = ['ID', 'Caso']
for _, _, n, _, _ in G:
    lab += [f'{n}\ncodice', f'{n}\nricalcolo', 'diff.']
c_base = 3 + 3 * len(G)  # prima colonna dei dati di base
lab += ['Dati di base:\nn. celle del bacino', 'z_max (m)', 'z_min (m)']
intest(ws, 1, lab, [6, 30] + [13, 13, 9] * len(G) + [14, 11, 11])
cN, cZmax, cZmin = COL(c_base), COL(c_base + 1), COL(c_base + 2)
POS = {}  # posizione delle colonne ricalcolo per i collegamenti
for r, i in enumerate(ids, 2):
    cella(ws, r, 1, i)
    cella(ws, r, 2, CASO.get(i, ''))
    cella(ws, r, c_base, int(V.loc[i, 'A_celle']), '#,##0', FBLU)
    cella(ws, r, c_base + 1, float(V.loc[i, 'Quota_max_m']), '0.00', FBLU)
    cella(ws, r, c_base + 2, float(V.loc[i, 'Quota_min_m']), '0.00', FBLU)
    for k, (cu, cv, n, fmt, tol) in enumerate(G):
        c0 = 3 + 3 * k
        cella(ws, r, c0, float(U.loc[i, cu]), fmt, FBLU)
        if cv == 'A_km2':
            ric = f'={cN}{r}*25/10^6'
        elif cv == 'Relief_m':
            ric = f'={cZmax}{r}-{cZmin}{r}'
        else:
            ric = float(V.loc[i, cv])
        cella(ws, r, c0 + 1, ric, fmt, F if isinstance(ric, str) else FBLU)
        L1, L2, Ld = COL(c0), COL(c0 + 1), COL(c0 + 2)
        cella(ws, r, c0 + 2, f'={L2}{r}-{L1}{r}', fmt)
        ws.conditional_formatting.add(f'{Ld}{r}', FormulaRule(formula=[f'ABS({Ld}{r})<={tol * 1.0001}'], fill=VERDE))
        ws.conditional_formatting.add(f'{Ld}{r}', FormulaRule(formula=[f'ABS({Ld}{r})>{tol * 1.0001}'], fill=ROSSO))
        POS[cv] = L2
last = N + 1
rt = last + 2
cella(ws, rt, 2, 'Ponti entro tolleranza (su ' + str(N) + ')', font=FB)
cella(ws, rt + 1, 2, 'Tolleranza (arrotondamento del codice)', font=FB)
for k, (cu, cv, n, fmt, tol) in enumerate(G):
    Ld = COL(3 + 3 * k + 2)
    cella(ws, rt + 1, 3 + 3 * k + 2, tol * 1.0001, '0.0E+00', FBLU)
    cella(ws, rt, 3 + 3 * k + 2, f'=SUMPRODUCT(--(ABS({Ld}2:{Ld}{last})<={Ld}{rt + 1}))', '0', FB)
ws.freeze_panes = 'C2'
ws.cell(rt + 3, 1, 'Verde: differenza entro l\'arrotondamento con cui il codice scrive il valore (A e L a 3 decimali, quote '
                   'e CN a 2, pendenze a 5). Ponte 167: ricalcolo in precisione singola (limite di memoria).').font = F

# ---------------- Parametri derivati ----------------
ws = wb.create_sheet('Parametri derivati')
inp = [('A (km²)', 'A_km2', None), ('L (km)', 'L_km', None), ('R (m)', 'Relief_m', None), ('AE (m)', 'AE_m', None),
       ('z_min (m)', None, cZmin), ('P (km)', 'P_km', None)]
der = [('FF', 'Form_factor_FF', '=B{r}/C{r}^2'),
       ('Melton', 'Melton_relative_relief', '=(D{r}/1000)/SQRT(B{r})'),
       ('HI', 'Hypsometric_integral', '=(E{r}-F{r})/D{r}'),
       ('Re', 'Elongation_ratio', '=2*SQRT(B{r}/PI())/C{r}'),
       ('Shape factor', 'Shape_factor', '=C{r}/SQRT(B{r})'),
       ('Kc', 'Compactness', '=G{r}/(2*SQRT(PI()*B{r}))'),
       ('Rc', 'Circularity', '=4*PI()*B{r}/G{r}^2'),
       ('Rr', 'Relief_Ratio', '=(D{r}/1000)/C{r}'),
       ('Rilievo relativo', 'Relative_relief', '=D{r}/(G{r}*1000)')]
lab = ['ID'] + [x[0] for x in inp]
for n, _, _ in der:
    lab += [f'{n}\nformula', f'{n}\ncodice', 'diff.']
intest(ws, 1, lab, [6] + [11] * 6 + [11, 11, 9] * len(der))
for r, i in enumerate(ids, 2):
    cella(ws, r, 1, i)
    for k, (nome, cv, colb) in enumerate(inp, 2):
        ref = f'={SP}!{POS[cv]}{r}' if cv else f'={SP}!{colb}{r}'
        cella(ws, r, k, ref, '0.000000', FVERDE, GIALLO)
    for k, (n, cu, f) in enumerate(der):
        c0 = 8 + 3 * k
        a = cella(ws, r, c0, f.format(r=r), '0.0000')
        b = cella(ws, r, c0 + 1, float(U.loc[i, cu]), '0.0000', FBLU)
        Ld = COL(c0 + 2)
        cella(ws, r, c0 + 2, f'={a.column_letter}{r}-{b.column_letter}{r}', '0.0000')
        ws.conditional_formatting.add(f'{Ld}{r}', FormulaRule(formula=[f'ABS({Ld}{r})<=0.0006'], fill=VERDE))
        ws.conditional_formatting.add(f'{Ld}{r}', FormulaRule(formula=[f'ABS({Ld}{r})>0.0006'], fill=ROSSO))
rt = last + 2
cella(ws, rt, 1, f'Ponti entro tolleranza (su {N})', font=FB)
for k in range(len(der)):
    Ld = COL(8 + 3 * k + 2)
    cella(ws, rt, 8 + 3 * k + 2, f'=SUMPRODUCT(--(ABS({Ld}2:{Ld}{last})<=0.0006))', '0', FB)
ws.freeze_panes = 'B2'
ws.cell(rt + 2, 1, 'Celle gialle (testo verde): grandezze collegate al foglio "Grandezze primarie", a precisione piena; '
                   'si possono sovrascrivere a mano. Verde: |diff.| ≤ 0,0006 (il codice arrotonda gli indici a 4 decimali).').font = F

# ---------------- Metodo indipendente ----------------
ws = wb.create_sheet('Metodo indipendente')
lab = ['ID', 'Caso', 'Area DB (km²)', 'Area A codice (km²)', 'Area B indipendente (km²)', 'Area comune A∩B (km²)',
       'Diff. B/A', 'IoU', 'Solo in A (km²)', 'Solo in B (km²)', 'Distanza sezioni A-B (m)', 'Regola sezione B',
       'L A (km)', 'L B (km)', 'AE A (m)', 'AE B (m)', 'CN A', 'CN B', 'Diff. A vs DB', 'Diff. B vs DB']
intest(ws, 1, lab, [6, 30, 11, 11, 11, 11, 9, 8, 10, 10, 11, 34, 9, 9, 9, 9, 8, 8, 10, 10])
for r, i in enumerate(ids, 2):
    v = V.loc[i]
    adb = pd.to_numeric(db.loc[i, 'SUPERFICIE BACINO (km2)'], errors='coerce')
    inter = float(v.A_km2 - v.Solo_A_km2)
    nn = lambda x: None if pd.isna(x) else float(x)
    row = [(i, None, F), (CASO.get(i, ''), None, F), (nn(adb), '0.000', FBLU),
           (f"={SP}!{POS['A_km2']}{r}", '0.000', FVERDE), (float(v.A_B_km2), '0.000', FBLU), (inter, '0.000', FBLU),
           (f'=E{r}/D{r}-1', '0.0%', F), (f'=F{r}/(D{r}+E{r}-F{r})', '0.0%', F), (f'=D{r}-F{r}', '0.000', F),
           (f'=E{r}-F{r}', '0.000', F), (float(v.Dist_sezioni_AB_m), '0.0', FBLU), (v.Regola_sezione_B, None, FBLU),
           (f"={SP}!{POS['L_km']}{r}", '0.000', FVERDE), (nn(v.L_B_km), '0.000', FBLU),
           (f"={SP}!{POS['AE_m']}{r}", '0.0', FVERDE), (float(v.AE_B_m), '0.0', FBLU),
           (f"={SP}!{POS['CN_medio']}{r}", '0.0', FVERDE), (float(v.CN_B), '0.0', FBLU),
           (f'=IF(C{r}="","",D{r}/C{r}-1)', '0.0%', F), (f'=IF(C{r}="","",E{r}/C{r}-1)', '0.0%', F)]
    for k, (val, fmt, font) in enumerate(row, 1):
        cella(ws, r, k, val, fmt, font)
    ws.conditional_formatting.add(f'H{r}', FormulaRule(formula=[f'H{r}>=0.90'], fill=VERDE))
    ws.conditional_formatting.add(f'H{r}', FormulaRule(formula=[f'H{r}<0.90'], fill=ROSSO))
rt = last + 2
cella(ws, rt, 2, 'Bacini concordanti (IoU ≥ 90%)', font=FB)
cella(ws, rt, 8, f'=COUNTIF(H2:H{last},">=0.9")', '0', FB)
ws.freeze_panes = 'C2'
wb.save(out)
print('salvato', out, N, 'ponti')
