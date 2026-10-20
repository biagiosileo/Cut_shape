"""
Consegna della revisione dei ponti fuori tolleranza: Excel con i tre gruppi (con colonna Conferma da
compilare), KML per Google Earth e ponti_proposti.csv (da usare solo dopo la conferma).

Uso: python revisione_consegna.py <revisione.csv> <revisione_non_compatibili.gpkg>
        <controllo_coordinate.gpkg> <ponti.csv> <out_dir>
"""
import os
import sys
from xml.sax.saxutils import escape

import geopandas as gpd
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

rcsv, gp, cc, pcsv, out = sys.argv[1:6]
R = pd.read_csv(rcsv)
E = gpd.read_file(cc, layer='esito_aree').set_index('ID')
L = gpd.read_file(cc, layer='db_attraversamento_inizio_fine').set_index('ID')


def ipotesi(r):
    t = []
    if 'ferrovia' in str(r['Nome opera']).lower() or 'sovrap' in str(r['Nome opera']).lower():
        t.append("Il nome indica un viadotto su ferrovia/strada: verificare che sia un attraversamento del reticolo.")
    if pd.notna(r.get('Dist_compat_2km_m')):
        t.append(f"Esiste una sezione con l'area del database a {r.Dist_compat_2km_m:.0f} m: le coordinate potrebbero "
                 f"indicare un altro attraversamento.")
    elif r.A_db_km2 > 3 * r.A_max_300m_km2:
        t.append(f"Nessun corso entro 300 m raggiunge l'area del database (massimo {r.A_max_300m_km2:.3g} km²) e "
                 f"nessuna sezione compatibile entro 2 km: area probabilmente riferita a un altro corso o sezione.")
    elif r.A_db_km2 < r.A_attuale_km2 / 3:
        t.append("Area del database molto piu' piccola di quella calcolata e nessuna sezione compatibile entro 2 km: "
                 "possibile fosso laterale non nel reticolo, oppure errore di trascrizione (unita' o virgola).")
    else:
        t.append("Nessuna sezione compatibile entro 2 km: verificare coordinate e valore di area.")
    return ' '.join(t)


def domanda(r):
    return (f"Con quale metodo e da quale punto di chiusura e' stata calcolata l'area di {r.A_db_km2:g} km² "
            f"({r.SP} km {r.km})? Le coordinate inizio/fine sono quelle dell'attraversamento del corso d'acqua?")


S = R[R.Gruppo == 'sposta'].copy()
F = R[R.Gruppo == 'soglia'].copy()
D = R[R.Gruppo == 'db'].copy()
D['Dist_inizio_fine_DB_m'] = D.ID.map(L['Dist_inizio_fine_m']).round(0)
D['Dist_punto_KML_m'] = D.ID.map(E['Dist_input_KML_m']).round(0)
D['Segnalazioni'] = D.ID.map(E['Segnalazioni'])
D['Ipotesi'] = [ipotesi(r) for _, r in D.iterrows()]
D['Domanda per il compilatore'] = [domanda(r) for _, r in D.iterrows()]

col_s = ['ID', 'Nome opera', 'SP', 'km', 'A_db_km2', 'A_attuale_km2', 'A_proposta_km2', 'Diff attuale', 'Diff proposta',
         'Dist_spostamento_m', 'Dist_dal_tracciato_m', 'Corso', 'Verifica', 'Conferma (SI/NO)', 'Note']
col_f = ['ID', 'Nome opera', 'SP', 'km', 'A_db_km2', 'A_attuale_km2', 'Soglia_celle', 'A_soglia_ridotta_km2',
         'A_proposta_km2', 'Diff attuale', 'Diff proposta', 'Dist_spostamento_m', 'Dist_dal_tracciato_m', 'Corso',
         'Verifica', 'Conferma (SI/NO)', 'Note']
col_d = ['ID', 'Nome opera', 'SP', 'km', 'A_db_km2', 'A_attuale_km2', 'Rapporto DB / calcolata', 'A_max_300m_km2',
         'Dist_compat_2km_m', 'A_compat_km2', 'Dist_inizio_fine_DB_m', 'Dist_punto_KML_m', 'Segnalazioni', 'Ipotesi',
         'Domanda per il compilatore']
for T, cols in ((S, col_s), (F, col_f), (D, col_d)):
    for c in cols:
        if c not in T:
            T[c] = None
xl = os.path.join(out, 'Revisione_non_compatibili.xlsx')
leggi = pd.DataFrame({'Revisione dei ponti fuori tolleranza (stesso DEM e stessa catena del codice)': [
    "Foglio 'Da spostare': sezione con area compatibile (+-15%) entro 300 m. Proposta = sezione compatibile piu' "
    "vicina al tracciato del ponte. Verificare su ortofoto (KML in Google Earth o GeoPackage in QGIS) e compilare "
    "la colonna Conferma.",
    "Foglio 'Fossi piccoli': area del database sotto la soglia del reticolo (0,1 km2). Calcolo con soglia ridotta "
    "solo per il ponte (meta' dell'area del database). Stessa verifica e colonna Conferma.",
    "Foglio 'Verifica database': nessuna sezione compatibile entro 300 m. Area massima entro 300 m, sezione "
    "compatibile entro 2 km, ipotesi e domanda da girare a chi ha compilato il database.",
    "Diff = area calcolata / area del database - 1 (formule). Verifica 'ok' = bacino ricalcolato dal punto proposto "
    "entro +-15%.",
    "ponti_proposti.csv: ponti.csv con le posizioni proposte; usarlo solo per i ponti confermati."]})
with pd.ExcelWriter(xl, engine='openpyxl') as w:
    leggi.to_excel(w, sheet_name='Leggimi', index=False)
    S[col_s].to_excel(w, sheet_name='Da spostare', index=False)
    F[col_f].to_excel(w, sheet_name='Fossi piccoli', index=False)
    D[col_d].to_excel(w, sheet_name='Verifica database', index=False)
wb = load_workbook(xl)
for nome, cols in (('Da spostare', col_s), ('Fossi piccoli', col_f), ('Verifica database', col_d)):
    ws = wb[nome]
    L_ = {c: chr(65 + k) for k, c in enumerate(cols)}
    for r in range(2, ws.max_row + 1):
        if 'Diff attuale' in L_:
            ws[f"{L_['Diff attuale']}{r}"] = f"={L_['A_attuale_km2']}{r}/{L_['A_db_km2']}{r}-1"
            ws[f"{L_['Diff proposta']}{r}"] = f"=IF({L_['A_proposta_km2']}{r}=\"\",\"\",{L_['A_proposta_km2']}{r}/{L_['A_db_km2']}{r}-1)"
            for c in ('Diff attuale', 'Diff proposta'):
                ws[f'{L_[c]}{r}'].number_format = '0.0%'
        if 'Rapporto DB / calcolata' in L_:
            ws[f"{L_['Rapporto DB / calcolata']}{r}"] = f"={L_['A_db_km2']}{r}/{L_['A_attuale_km2']}{r}"
            ws[f"{L_['Rapporto DB / calcolata']}{r}"].number_format = '0.00'
        for c in ws[r]:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    if 'Conferma (SI/NO)' in L_:
        dv = DataValidation(type='list', formula1='"SI,NO"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{L_['Conferma (SI/NO)']}2:{L_['Conferma (SI/NO)']}{ws.max_row}")
        for r in range(2, ws.max_row + 1):
            ws[f"{L_['Conferma (SI/NO)']}{r}"].fill = PatternFill('solid', fgColor='FFF2CC')
    for k, c in enumerate(cols):
        ws.column_dimensions[chr(65 + k)].width = 60 if c in ('Ipotesi', 'Domanda per il compilatore') else (
            28 if c in ('Nome opera', 'Segnalazioni', 'Note', 'Corso') else 12)
    ws.freeze_panes = 'C2'
for ws in wb.worksheets:
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDDDDD')
        c.alignment = Alignment(wrap_text=True, vertical='center')
wb['Leggimi'].column_dimensions['A'].width = 140
wb.save(xl)

# ponti.csv con le posizioni proposte (solo gruppi sposta e soglia)
P = pd.read_csv(pcsv)
nuovi = R.dropna(subset=['X_nuovo']).set_index('ID')
k = P.id.isin(nuovi.index)
P.loc[k, 'X'] = P.loc[k, 'id'].map(nuovi.X_nuovo)
P.loc[k, 'Y'] = P.loc[k, 'id'].map(nuovi.Y_nuovo)
P.to_csv(os.path.join(out, 'ponti_proposti.csv'), index=False)

# KML per Google Earth
STILI = {'punto_attuale': ('ff0000ff', 'attuale'), 'punto_proposto': ('ff00ff00', 'proposto'),
         'spostamento': ('ff00ffff', 'spostamento'), 'bacino_attuale': ('ff0000ff', 'bacino attuale'),
         'bacino_proposto': ('ff00ff00', 'bacino proposto'), 'tracciato_ponte': ('ffff00ff', 'tracciato ponte')}


def coords(g):
    return ' '.join(f'{x:.7f},{y:.7f},0' for x, y in g.coords)


def geom_kml(g):
    if g.geom_type == 'Point':
        return f'<Point><coordinates>{g.x:.7f},{g.y:.7f},0</coordinates></Point>'
    if g.geom_type == 'LineString':
        return f'<LineString><coordinates>{coords(g)}</coordinates></LineString>'
    if g.geom_type == 'Polygon':
        inn = ''.join(f'<innerBoundaryIs><LinearRing><coordinates>{coords(i)}</coordinates></LinearRing></innerBoundaryIs>'
                      for i in g.interiors)
        return (f'<Polygon><outerBoundaryIs><LinearRing><coordinates>{coords(g.exterior)}</coordinates></LinearRing>'
                f'</outerBoundaryIs>{inn}</Polygon>')
    return '<MultiGeometry>' + ''.join(geom_kml(p) for p in g.geoms) + '</MultiGeometry>'


k = ['<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document>'
     '<name>Revisione ponti fuori tolleranza</name>']
for s, (col, _) in STILI.items():
    k.append(f'<Style id="{s}"><IconStyle><color>{col}</color><scale>0.9</scale></IconStyle>'
             f'<LineStyle><color>{col}</color><width>2</width></LineStyle><PolyStyle><fill>0</fill></PolyStyle></Style>')
for s, (_, titolo) in STILI.items():
    g = gpd.read_file(gp, layer=s).to_crs(4326)
    if s.startswith('bacino'):
        g['geometry'] = g.geometry.simplify(0.00003)
    k.append(f'<Folder><name>{escape(titolo)}</name>')
    for _, r in g.iterrows():
        info = ', '.join(f'{c}={r[c]}' for c in g.columns if c not in ('geometry',) and pd.notna(r[c]) and r[c] != '')
        k.append(f'<Placemark><name>{int(r.ID)}</name><description>{escape(info)}</description>'
                 f'<styleUrl>#{s}</styleUrl>{geom_kml(r.geometry)}</Placemark>')
    k.append('</Folder>')
k.append('</Document></kml>')
open(os.path.join(out, 'Revisione_non_compatibili.kml'), 'w', encoding='utf-8').write(''.join(k))
print('ok')
